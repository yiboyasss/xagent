from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from threading import Event, Lock, local
from uuid import UUID

import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import QueuePool

from xagent.core.agent.service import AgentService
from xagent.core.file_storage.factory import (
    get_unscoped_file_storage,
    get_user_file_storage,
)
from xagent.core.file_storage.storage import FsspecFileStorage
from xagent.core.tools.adapters.vibe.factory import ToolFactory
from xagent.core.workspace import TaskWorkspace, WorkspaceManager
from xagent.web.models import Base
from xagent.web.models.task import Task
from xagent.web.models.uploaded_file import UploadedFile
from xagent.web.models.user import User


@pytest.fixture
def constrained_workspace_db():
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=QueuePool,
        pool_size=1,
        max_overflow=0,
        pool_timeout=0.1,
    )
    SessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)
    Base.metadata.create_all(bind=engine)
    db = SessionLocal()
    try:
        yield engine, SessionLocal, db
    finally:
        db.close()
        engine.dispose()


def _seed_workspace_task(db, *, task_id: int, username: str) -> User:
    user = User(username=username, password_hash="hash")
    db.add(user)
    db.flush()
    db.add(Task(id=task_id, user_id=user.id, title="Workspace task"))
    db.commit()
    return user


def test_inline_delivery_rejects_registration_without_durable_task_owner(
    monkeypatch, tmp_path, constrained_workspace_db
):
    from xagent.core.inline_file_delivery import InlineFileDelivery

    engine, SessionLocal, db = constrained_workspace_db
    monkeypatch.setattr("xagent.web.models.database.get_session_local", lambda: SessionLocal)
    monkeypatch.setattr("xagent.core.storage.manager.create_db_session", SessionLocal)
    workspace = TaskWorkspace(id="web_task_9999", base_dir=str(tmp_path))
    delivery = InlineFileDelivery(workspace)
    result = delivery.transform("[report.csv](data:text/csv;base64,YSwK)")
    assert result == "report.csv (attachment unavailable)"
    assert db.query(UploadedFile).count() == 0
    db.rollback()
    assert engine.pool.checkedout() == 0
    assert not list(workspace.output_dir.iterdir())


@pytest.mark.asyncio
@pytest.mark.parametrize("file_kind", ["csv", "xlsx", "png"])
@pytest.mark.parametrize("delivery_kind", ["link", "fence"])
async def test_inline_delivery_uses_task_owned_durable_file_registration(
    monkeypatch, tmp_path, constrained_workspace_db, file_kind, delivery_kind
):
    import base64
    from io import BytesIO
    from pathlib import Path

    from fastapi import FastAPI
    from httpx import ASGITransport, AsyncClient
    from openpyxl import Workbook, load_workbook

    from xagent.core.agent import PatternRuntime
    from xagent.core.inline_file_delivery import InlineFileDelivery
    from xagent.web.api.auth import create_access_token
    from xagent.web.api.files import file_router
    from xagent.web.models.database import get_db

    engine, SessionLocal, db = constrained_workspace_db
    user = _seed_workspace_task(db, task_id=9019, username="inline-delivery-owner")
    user_id = int(user.id)
    owner_headers = {
        "Authorization": "Bearer "
        + create_access_token(data={"sub": user.username, "user_id": user_id})
    }
    stranger = _seed_workspace_task(db, task_id=9020, username="inline-delivery-other")
    other_headers = {
        "Authorization": "Bearer "
        + create_access_token(data={"sub": stranger.username, "user_id": stranger.id})
    }
    db.rollback()
    monkeypatch.setenv("XAGENT_FILE_STORAGE_URI", (tmp_path / "objects").as_uri())
    monkeypatch.setenv("XAGENT_UPLOADS_DIR", str(tmp_path / "uploads"))
    monkeypatch.setenv("XAGENT_FILE_DELIVERY_REDIRECT_ENABLED", "false")
    get_unscoped_file_storage.cache_clear()
    monkeypatch.setattr("xagent.web.models.database.get_session_local", lambda: SessionLocal)
    monkeypatch.setattr(
        "xagent.web.services.uploaded_file_store.get_session_local",
        lambda: SessionLocal,
    )
    monkeypatch.setattr("xagent.core.storage.manager.create_db_session", SessionLocal)
    workspace = TaskWorkspace(
        id="web_task_9019", base_dir=str(tmp_path / "uploads" / f"user_{user_id}")
    )

    rows = [
        ["order_id", "reason_code", "remake_rate"],
        ["A001", "CUST", 0.04],
        ["A002", "POWER", 0.15],
        ["A003", "CUST", 0.07],
        ["A004", "COATING", 0.11],
    ]
    if file_kind == "xlsx":
        workbook = Workbook()
        sheet = workbook.active
        for row in rows:
            sheet.append(row)
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        data = output.getvalue()
        mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    elif file_kind == "png":
        data = base64.b64decode(
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8"
            "/x8AAwMCAO+aXioAAAAASUVORK5CYII="
        )
        mime = "image/png"
    else:
        # The four-row replacement used for the original file-delivery smoke
        # test, not the unavailable original TC5 business attachment.
        data = "\r\n".join(",".join(map(str, row)) for row in rows).encode() + b"\r\n"
        mime = "text/csv"
    name = f"remake_analysis.{file_kind}"
    encoded = base64.b64encode(data).decode()
    if delivery_kind == "fence":
        source = f"```base64 filename={name}\n{encoded}\n```"
    else:
        source = f"[{name}](data:{mime};base64,{encoded})"
        if file_kind == "png":
            source = "!" + source

    def isolated_db():
        with SessionLocal() as session:
            yield session

    app = FastAPI()
    app.include_router(file_router)
    app.dependency_overrides[get_db] = isolated_db
    try:
        runtime = PatternRuntime(inline_file_delivery=InlineFileDelivery(workspace))
        answer = await runtime.prepare_final_answer(source)
        assert "base64" not in answer
        with SessionLocal() as verify_db:
            record = verify_db.query(UploadedFile).one()
            assert record.user_id == user_id
            assert record.task_id == 9019
            assert record.storage_status == "available"
            prefix = "!" if file_kind == "png" and delivery_kind == "link" else ""
            assert answer == f"{prefix}[{name}](file:{record.file_id})"
            file_id = record.file_id
            local_path = Path(record.storage_path)
            assert local_path.read_bytes() == data
        assert engine.pool.checkedout() == 0
        # Exercise production HTTP routes and real Bearer authentication, not
        # a mocked FileRef or a direct materialize() call. Repeat without the
        # workspace copy to cover durable-storage recovery on both endpoints.
        async with AsyncClient(
            transport=ASGITransport(app=app), base_url="http://testserver"
        ) as client:
            for remove_local in (False, True):
                for endpoint in ("preview", "download"):
                    if remove_local:
                        local_path.unlink(missing_ok=True)
                    url = f"/api/files/{endpoint}/{file_id}"
                    response = await client.get(url, headers=owner_headers)
                    assert response.status_code == 200, response.text
                    assert response.content == data
                    assert response.headers["content-type"].split(";")[0] == mime
                    assert response.headers["x-content-type-options"] == "nosniff"
                    if endpoint == "preview":
                        assert response.headers["content-disposition"] == "inline"
                    else:
                        assert name in response.headers["content-disposition"]
                    if file_kind == "xlsx":
                        downloaded = load_workbook(BytesIO(response.content))
                        assert list(downloaded.active.values) == list(map(tuple, rows))
                        downloaded.close()
                    assert (await client.get(url)).status_code in (401, 403)
                    assert (await client.get(url, headers=other_headers)).status_code == 403
            assert engine.pool.checkedout() == 0
    finally:
        get_unscoped_file_storage.cache_clear()


def _assert_task_output_generation_key(
    storage_key: str,
    *,
    user_id: int,
    task_id: int,
    file_id: str,
    relative_path: str,
) -> None:
    prefix = f"users/{user_id}/tasks/{task_id}/outputs/{file_id}/_versions/"
    assert storage_key.startswith(prefix)
    generation_and_path = storage_key.removeprefix(prefix)
    generation, actual_relative_path = generation_and_path.split("/", 1)
    assert UUID(generation)
    assert actual_relative_path == relative_path


def test_workspace_register_create_releases_pool_before_durable_put(
    monkeypatch,
    tmp_path,
    mock_workspace_db,
    constrained_workspace_db,
):
    del mock_workspace_db
    engine, _SessionLocal, db = constrained_workspace_db
    user = _seed_workspace_task(
        db,
        task_id=9001,
        username="workspace-create-boundary-user",
    )
    object_root = tmp_path / "objects"
    monkeypatch.setenv("XAGENT_FILE_STORAGE_URI", object_root.as_uri())
    get_unscoped_file_storage.cache_clear()

    observed_checked_out = []
    original_put_file = FsspecFileStorage.put_file

    def put_file_spy(self, source, key, content_type=None):
        observed_checked_out.append(engine.pool.checkedout())
        return original_put_file(self, source, key, content_type)

    monkeypatch.setattr(FsspecFileStorage, "put_file", put_file_spy)

    workspace = TaskWorkspace(
        id="web_task_9001",
        base_dir=str(tmp_path / "workspaces"),
    )
    output_path = workspace.output_dir / "report.txt"
    output_path.write_text("new output", encoding="utf-8")

    file_id = workspace.register_file(str(output_path), db_session=db)
    db.commit()

    assert observed_checked_out == [0]
    record = db.query(UploadedFile).filter(UploadedFile.file_id == file_id).one()
    assert record.user_id == user.id


def test_workspace_register_resync_releases_pool_before_durable_put(
    monkeypatch,
    tmp_path,
    mock_workspace_db,
    constrained_workspace_db,
):
    del mock_workspace_db
    engine, SessionLocal, db = constrained_workspace_db
    _seed_workspace_task(
        db,
        task_id=9002,
        username="workspace-resync-boundary-user",
    )
    object_root = tmp_path / "objects"
    monkeypatch.setenv("XAGENT_FILE_STORAGE_URI", object_root.as_uri())
    get_unscoped_file_storage.cache_clear()

    workspace = TaskWorkspace(
        id="web_task_9002",
        base_dir=str(tmp_path / "workspaces"),
    )
    output_path = workspace.output_dir / "report.txt"
    output_path.write_text("first generation", encoding="utf-8")
    file_id = workspace.register_file(str(output_path), db_session=db)
    db.commit()

    observed_checked_out = []
    original_put_file = FsspecFileStorage.put_file

    def put_file_spy(self, source, key, content_type=None):
        observed_checked_out.append(engine.pool.checkedout())
        return original_put_file(self, source, key, content_type)

    monkeypatch.setattr(FsspecFileStorage, "put_file", put_file_spy)
    output_path.write_text("second generation", encoding="utf-8")

    assert workspace.register_file(str(output_path), db_session=db) == file_id
    db.commit()

    assert observed_checked_out == [0]


def test_workspace_registration_commit_survives_caller_rollback(
    monkeypatch,
    tmp_path,
    mock_workspace_db,
    constrained_workspace_db,
):
    del mock_workspace_db
    engine, SessionLocal, caller_db = constrained_workspace_db
    user = _seed_workspace_task(
        caller_db,
        task_id=9006,
        username="workspace-owner-commit-user",
    )
    user_id = int(user.id)
    object_root = tmp_path / "objects"
    monkeypatch.setenv("XAGENT_FILE_STORAGE_URI", object_root.as_uri())
    get_unscoped_file_storage.cache_clear()

    from xagent.web.services import uploaded_file_store as store_module

    monkeypatch.setattr(store_module, "get_session_local", lambda: SessionLocal)

    workspace = TaskWorkspace(
        id="web_task_9006",
        base_dir=str(tmp_path / "workspaces"),
    )
    output_path = workspace.output_dir / "report.txt"
    output_path.write_text("independently committed", encoding="utf-8")

    file_id = workspace.register_file(str(output_path), db_session=caller_db)
    caller_db.rollback()

    assert engine.pool.checkedout() == 0
    with SessionLocal() as verify_db:
        record = verify_db.query(UploadedFile).filter(UploadedFile.file_id == file_id).one()
        assert record.user_id == user_id
        storage_key = str(record.storage_key)
    assert get_user_file_storage(user_id).exists(storage_key) is True


def test_workspace_commit_ack_loss_keeps_referenced_generation(
    monkeypatch,
    tmp_path,
    mock_workspace_db,
    constrained_workspace_db,
):
    del mock_workspace_db
    engine, SessionLocal, caller_db = constrained_workspace_db
    user = _seed_workspace_task(
        caller_db,
        task_id=9007,
        username="workspace-commit-ack-user",
    )
    user_id = int(user.id)
    object_root = tmp_path / "objects"
    monkeypatch.setenv("XAGENT_FILE_STORAGE_URI", object_root.as_uri())
    get_unscoped_file_storage.cache_clear()

    from xagent.web.services import uploaded_file_store as store_module

    monkeypatch.setattr(store_module, "get_session_local", lambda: SessionLocal)

    read_db = SessionLocal()
    write_db = SessionLocal()
    real_commit = write_db.commit

    def commit_then_lose_ack():
        real_commit()
        raise RuntimeError("metadata commit acknowledgement lost")

    monkeypatch.setattr(write_db, "commit", commit_then_lose_ack)
    registration_sessions = iter((read_db, write_db))

    workspace = TaskWorkspace(
        id="web_task_9007",
        base_dir=str(tmp_path / "workspaces"),
    )
    monkeypatch.setattr(
        workspace,
        "_create_registration_session",
        lambda _reference_db: next(registration_sessions),
        raising=False,
    )
    output_path = workspace.output_dir / "report.txt"
    output_path.write_text("committed before ack loss", encoding="utf-8")

    with pytest.raises(
        RuntimeError,
        match="metadata commit acknowledgement lost",
    ):
        workspace.register_file(str(output_path), db_session=caller_db)

    assert engine.pool.checkedout() == 0
    with SessionLocal() as verify_db:
        record = verify_db.query(UploadedFile).one()
        assert record.user_id == user_id
        storage_key = str(record.storage_key)
    assert get_user_file_storage(user_id).exists(storage_key) is True


def test_workspace_delegated_registration_releases_pool_before_canonical_copy(
    monkeypatch,
    tmp_path,
    mock_workspace_db,
    constrained_workspace_db,
):
    del mock_workspace_db
    engine, _SessionLocal, db = constrained_workspace_db
    user = _seed_workspace_task(
        db,
        task_id=9003,
        username="workspace-copy-boundary-user",
    )
    object_root = tmp_path / "objects"
    monkeypatch.setenv("XAGENT_FILE_STORAGE_URI", object_root.as_uri())
    get_unscoped_file_storage.cache_clear()

    import xagent.core.workspace as workspace_module

    observed_checked_out = []
    original_copy2 = workspace_module.shutil.copy2

    def copy2_spy(source, target):
        observed_checked_out.append(engine.pool.checkedout())
        return original_copy2(source, target)

    monkeypatch.setattr(workspace_module.shutil, "copy2", copy2_spy)

    workspace = TaskWorkspace(
        id="agent_2_boundary",
        base_dir=str(tmp_path / "workspaces"),
        db_task_id=9003,
    )
    output_path = workspace.output_dir / "report.txt"
    output_path.write_text("delegated output", encoding="utf-8")

    file_id = workspace.register_file(str(output_path), db_session=db)
    db.commit()

    assert observed_checked_out == [0]
    record = db.query(UploadedFile).filter(UploadedFile.file_id == file_id).one()
    assert record.storage_path == str(
        tmp_path / "workspaces" / f"user_{user.id}" / "web_task_9003" / "output" / "report.txt"
    )


def test_workspace_register_compensates_staged_generation_on_metadata_failure(
    monkeypatch,
    tmp_path,
    mock_workspace_db,
    constrained_workspace_db,
):
    del mock_workspace_db
    engine, SessionLocal, db = constrained_workspace_db
    user = _seed_workspace_task(
        db,
        task_id=9004,
        username="workspace-compensation-user",
    )
    object_root = tmp_path / "objects"
    monkeypatch.setenv("XAGENT_FILE_STORAGE_URI", object_root.as_uri())
    get_unscoped_file_storage.cache_clear()

    from xagent.web.services import uploaded_file_store as store_module

    monkeypatch.setattr(store_module, "get_session_local", lambda: SessionLocal)
    staged_files = []
    original_stage = store_module.stage_uploaded_file_from_local_path

    def stage_spy(**kwargs):
        staged = original_stage(**kwargs)
        staged_files.append(staged)
        return staged

    def fail_metadata(self, staged, *, expected, allow_task_rebind=False):
        del self, staged, expected, allow_task_rebind
        raise store_module.UploadedFileVersionConflict("injected metadata failure")

    monkeypatch.setattr(store_module, "stage_uploaded_file_from_local_path", stage_spy)
    monkeypatch.setattr(
        store_module.UploadedFileStore,
        "upsert_already_durable",
        fail_metadata,
    )

    workspace = TaskWorkspace(
        id="web_task_9004",
        base_dir=str(tmp_path / "workspaces"),
    )
    output_path = workspace.output_dir / "report.txt"
    output_path.write_text("uncommitted generation", encoding="utf-8")

    with pytest.raises(
        store_module.UploadedFileVersionConflict,
        match="injected metadata failure",
    ):
        workspace.register_file(str(output_path), db_session=db)

    assert engine.pool.checkedout() == 0
    assert len(staged_files) == 1
    storage = get_user_file_storage(int(user.id))
    assert storage.exists(staged_files[0].storage_key) is False
    assert db.query(UploadedFile).count() == 0


def test_workspace_register_cas_collision_compensates_only_losing_generation(
    monkeypatch,
    tmp_path,
    mock_workspace_db,
    constrained_workspace_db,
):
    del mock_workspace_db
    engine, SessionLocal, db = constrained_workspace_db
    user = _seed_workspace_task(
        db,
        task_id=9005,
        username="workspace-cas-collision-user",
    )
    object_root = tmp_path / "objects"
    monkeypatch.setenv("XAGENT_FILE_STORAGE_URI", object_root.as_uri())
    get_unscoped_file_storage.cache_clear()

    workspace = TaskWorkspace(
        id="web_task_9005",
        base_dir=str(tmp_path / "workspaces"),
    )
    output_path = workspace.output_dir / "report.txt"
    output_path.write_text("original", encoding="utf-8")
    file_id = workspace.register_file(str(output_path), db_session=db)
    db.commit()

    from xagent.web.services import uploaded_file_store as store_module

    monkeypatch.setattr(store_module, "get_session_local", lambda: SessionLocal)
    staged_files = []
    winner_key = (
        f"users/{user.id}/tasks/9005/outputs/{file_id}/"
        "_versions/11111111111111111111111111111111/output/report.txt"
    )
    storage = get_user_file_storage(int(user.id))
    original_stage = store_module.stage_uploaded_file_from_local_path

    def stage_then_publish_winner(**kwargs):
        staged = original_stage(**kwargs)
        staged_files.append(staged)
        winner = storage.put_bytes(b"concurrent winner", winner_key, "text/plain")
        with SessionLocal() as winner_db:
            record = winner_db.query(UploadedFile).filter(UploadedFile.file_id == file_id).one()
            record.storage_backend = storage.backend
            record.storage_key = winner.key
            record.storage_uri = winner.uri
            record.checksum = winner.checksum
            record.etag = winner.etag
            record.file_size = len(b"concurrent winner")
            record.storage_status = "available"
            winner_db.commit()
        return staged

    monkeypatch.setattr(
        store_module,
        "stage_uploaded_file_from_local_path",
        stage_then_publish_winner,
    )
    output_path.write_text("losing writer", encoding="utf-8")

    with pytest.raises(store_module.UploadedFileVersionConflict):
        workspace.register_file(str(output_path), db_session=db)

    assert engine.pool.checkedout() == 0
    assert len(staged_files) == 1
    assert storage.exists(staged_files[0].storage_key) is False
    assert storage.exists(winner_key) is True
    current = db.query(UploadedFile).filter(UploadedFile.file_id == file_id).one()
    assert current.storage_key == winner_key
    assert current.checksum == storage.content_hash(winner_key)


def test_workspace_register_file_writes_durable_storage(monkeypatch, tmp_path, mock_workspace_db):
    # Override the global autouse fixture from tests/conftest.py for this module.
    del mock_workspace_db
    object_root = tmp_path / "objects"
    monkeypatch.setenv("XAGENT_FILE_STORAGE_URI", object_root.as_uri())
    get_unscoped_file_storage.cache_clear()

    engine = create_engine("sqlite:///:memory:", connect_args={"check_same_thread": False})
    SessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)
    Base.metadata.create_all(bind=engine)
    db = SessionLocal()
    try:
        user = User(username="workspace-user", password_hash="hash")
        db.add(user)
        db.flush()
        task = Task(id=123, user_id=user.id, title="Workspace task")
        db.add(task)
        db.commit()

        workspace = TaskWorkspace(id="web_task_123", base_dir=str(tmp_path / "workspaces"))
        output_path = workspace.output_dir / "report.txt"
        output_path.write_text("workspace output", encoding="utf-8")

        file_id = workspace.register_file(str(output_path), db_session=db)
        db.commit()

        record = db.query(UploadedFile).filter(UploadedFile.file_id == file_id).one()
        assert record.storage_path == str(output_path)
        assert record.storage_status == "available"
        assert record.storage_backend == "file"
        _assert_task_output_generation_key(
            str(record.storage_key),
            user_id=int(user.id),
            task_id=123,
            file_id=file_id,
            relative_path="output/report.txt",
        )
        assert record.workspace_relative_path == "output/report.txt"
        assert record.workspace_category == "output"

        object_files = [path for path in object_root.rglob("*") if path.is_file()]
        assert len(object_files) == 1
        assert object_files[0].read_text(encoding="utf-8") == "workspace output"
    finally:
        db.close()
        engine.dispose()


def test_agent_workspace_register_file_uses_explicit_db_task_id(
    monkeypatch, tmp_path, mock_workspace_db
):
    # Override the global autouse fixture from tests/conftest.py for this module.
    del mock_workspace_db
    object_root = tmp_path / "objects"
    monkeypatch.setenv("XAGENT_FILE_STORAGE_URI", object_root.as_uri())
    get_unscoped_file_storage.cache_clear()

    engine = create_engine("sqlite:///:memory:", connect_args={"check_same_thread": False})
    SessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)
    Base.metadata.create_all(bind=engine)
    db = SessionLocal()
    try:
        user = User(username="delegated-workspace-user", password_hash="hash")
        db.add(user)
        db.flush()
        task = Task(id=321, user_id=user.id, title="Parent task")
        db.add(task)
        db.commit()

        workspace = TaskWorkspace(
            id="agent_2_abcd1234",
            base_dir=str(tmp_path / "workspaces"),
            db_task_id=321,
        )
        output_path = workspace.output_dir / "report.txt"
        output_path.write_text("delegated output", encoding="utf-8")

        file_id = workspace.register_file(str(output_path), db_session=db)
        db.commit()

        record = db.query(UploadedFile).filter(UploadedFile.file_id == file_id).one()
        canonical_path = (
            tmp_path / "workspaces" / f"user_{user.id}" / "web_task_321" / "output" / "report.txt"
        )
        assert workspace.current_task_id == 321
        assert record.user_id == user.id
        assert record.task_id == 321
        assert record.storage_path == str(canonical_path)
        assert canonical_path.read_text(encoding="utf-8") == "delegated output"
        _assert_task_output_generation_key(
            str(record.storage_key),
            user_id=int(user.id),
            task_id=321,
            file_id=file_id,
            relative_path="output/report.txt",
        )
        assert record.workspace_relative_path == "output/report.txt"
        assert record.workspace_category == "output"
    finally:
        db.close()
        engine.dispose()


def test_agent_workspace_register_file_rebinds_existing_output_to_db_task_id(
    monkeypatch, tmp_path, mock_workspace_db
):
    # Override the global autouse fixture from tests/conftest.py for this module.
    del mock_workspace_db
    object_root = tmp_path / "objects"
    monkeypatch.setenv("XAGENT_FILE_STORAGE_URI", object_root.as_uri())
    get_unscoped_file_storage.cache_clear()

    engine = create_engine("sqlite:///:memory:", connect_args={"check_same_thread": False})
    SessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)
    Base.metadata.create_all(bind=engine)
    db = SessionLocal()
    try:
        user = User(username="delegated-rebind-user", password_hash="hash")
        db.add(user)
        db.flush()
        parent_task = Task(id=321, user_id=user.id, title="Parent task")
        worker_task = Task(id=322, user_id=user.id, title="Worker task")
        db.add_all([parent_task, worker_task])
        db.commit()

        workspace = TaskWorkspace(
            id="agent_2_abcd1234",
            base_dir=str(tmp_path / "workspaces"),
            db_task_id=321,
        )
        output_path = workspace.output_dir / "report.txt"
        output_path.write_text("worker output", encoding="utf-8")

        from xagent.web.services.uploaded_file_store import UploadedFileStore

        record = UploadedFileStore(db).create_from_local_path(
            local_path=output_path,
            user_id=int(user.id),
            file_id="worker-output",
            task_id=322,
            filename="report.txt",
            workspace_relative_path="output/report.txt",
            workspace_category="output",
            mime_type="text/plain",
        )
        db.commit()

        output_path.write_text("parent-visible output", encoding="utf-8")
        file_id = workspace.register_file(str(output_path), db_session=db)
        db.commit()

        assert file_id == "worker-output"
        db.refresh(record)
        canonical_path = (
            tmp_path / "workspaces" / f"user_{user.id}" / "web_task_321" / "output" / "report.txt"
        )
        assert record.user_id == user.id
        assert record.task_id == 321
        assert record.storage_path == str(canonical_path)
        assert canonical_path.read_text(encoding="utf-8") == "parent-visible output"
        _assert_task_output_generation_key(
            str(record.storage_key),
            user_id=int(user.id),
            task_id=321,
            file_id="worker-output",
            relative_path="output/report.txt",
        )
        assert record.workspace_relative_path == "output/report.txt"
        assert record.workspace_category == "output"
        assert record.file_size == len("parent-visible output")
    finally:
        db.close()
        engine.dispose()


def test_agent_workspace_register_file_avoids_parent_output_name_collision(
    monkeypatch, tmp_path, mock_workspace_db
):
    del mock_workspace_db
    object_root = tmp_path / "objects"
    monkeypatch.setenv("XAGENT_FILE_STORAGE_URI", object_root.as_uri())
    get_unscoped_file_storage.cache_clear()

    engine = create_engine("sqlite:///:memory:", connect_args={"check_same_thread": False})
    SessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)
    Base.metadata.create_all(bind=engine)
    db = SessionLocal()
    try:
        user = User(username="delegated-collision-user", password_hash="hash")
        db.add(user)
        db.flush()
        task = Task(id=321, user_id=user.id, title="Parent task")
        db.add(task)
        db.commit()

        parent_output_path = (
            tmp_path / "workspaces" / f"user_{user.id}" / "web_task_321" / "output" / "report.txt"
        )
        parent_output_path.parent.mkdir(parents=True)
        parent_output_path.write_text("existing parent output", encoding="utf-8")
        db.add(
            UploadedFile(
                file_id="existing-parent-output",
                user_id=user.id,
                task_id=321,
                filename="report.txt",
                storage_path=str(parent_output_path),
                mime_type="text/plain",
                file_size=len("existing parent output"),
                workspace_relative_path="output/report.txt",
                workspace_category="output",
            )
        )
        db.commit()

        workspace = TaskWorkspace(
            id="agent_2_abcd1234",
            base_dir=str(tmp_path / "workspaces"),
            db_task_id=321,
        )
        output_path = workspace.output_dir / "report.txt"
        output_path.write_text("delegated output", encoding="utf-8")

        file_id = workspace.register_file(str(output_path), db_session=db)
        db.commit()

        record = db.query(UploadedFile).filter(UploadedFile.file_id == file_id).one()
        canonical_path = parent_output_path.with_name("report_1.txt")
        assert parent_output_path.read_text(encoding="utf-8") == "existing parent output"
        assert canonical_path.read_text(encoding="utf-8") == "delegated output"
        assert record.storage_path == str(canonical_path)
        assert record.workspace_relative_path == "output/report_1.txt"
        _assert_task_output_generation_key(
            str(record.storage_key),
            user_id=int(user.id),
            task_id=321,
            file_id=file_id,
            relative_path="output/report_1.txt",
        )
    finally:
        db.close()
        engine.dispose()


def test_agent_workspace_register_file_is_idempotent_after_canonicalization(
    monkeypatch, tmp_path, mock_workspace_db
):
    del mock_workspace_db
    object_root = tmp_path / "objects"
    monkeypatch.setenv("XAGENT_FILE_STORAGE_URI", object_root.as_uri())
    get_unscoped_file_storage.cache_clear()

    engine = create_engine("sqlite:///:memory:", connect_args={"check_same_thread": False})
    SessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)
    Base.metadata.create_all(bind=engine)
    db = SessionLocal()
    try:
        user = User(username="delegated-idempotent-user", password_hash="hash")
        db.add(user)
        db.flush()
        task = Task(id=321, user_id=user.id, title="Parent task")
        db.add(task)
        db.commit()

        workspace = TaskWorkspace(
            id="agent_2_abcd1234",
            base_dir=str(tmp_path / "workspaces"),
            db_task_id=321,
        )
        output_path = workspace.output_dir / "report.txt"
        output_path.write_text("first delegated output", encoding="utf-8")

        file_id = workspace.register_file(str(output_path), db_session=db)
        db.commit()

        output_path.write_text("updated delegated output", encoding="utf-8")
        second_file_id = workspace.register_file(str(output_path), db_session=db)
        db.commit()

        assert second_file_id == file_id
        records = db.query(UploadedFile).all()
        assert len(records) == 1
        record = records[0]
        canonical_path = (
            tmp_path / "workspaces" / f"user_{user.id}" / "web_task_321" / "output" / "report.txt"
        )
        assert record.storage_path == str(canonical_path)
        assert record.workspace_relative_path == "output/report.txt"
        assert record.file_size == len("updated delegated output")
        assert canonical_path.read_text(encoding="utf-8") == "updated delegated output"

        assert (object_root / str(record.storage_key)).read_text(
            encoding="utf-8"
        ) == "updated delegated output"
    finally:
        db.close()
        engine.dispose()


@pytest.mark.parametrize(
    ("durable_segments", "expected"),
    [
        (None, ("tenant-a",)),
        ((), ()),
    ],
)
def test_tool_factory_workspace_preserves_db_task_id(tmp_path, durable_segments, expected):
    workspace = ToolFactory.create_workspace(
        {
            "base_dir": str(tmp_path / "workspaces"),
            "task_id": "agent_2_abcd1234",
            "db_task_id": 654,
            "__xagent_file_operation_access_version": 1,
            "user_id": 7,
            "scope_segments": ("tenant-a",),
            **({} if durable_segments is None else {"durable_storage_segments": durable_segments}),
        }
    )

    assert workspace is not None
    assert workspace.id == "agent_2_abcd1234"
    assert workspace.db_task_id == 654
    assert workspace.current_task_id == 654
    assert workspace.owner_user_id == 7
    assert workspace.file_operation_access_version == 1
    assert workspace.durable_storage_segments == expected


def test_workspace_manager_updates_cached_workspace_db_task_id(tmp_path):
    manager = WorkspaceManager()
    workspace = manager.get_or_create_workspace(
        str(tmp_path / "workspaces"),
        "agent_2_abcd1234",
    )
    assert workspace.db_task_id is None
    assert workspace.current_task_id is None

    same_workspace = manager.get_or_create_workspace(
        str(tmp_path / "workspaces"),
        "agent_2_abcd1234",
        db_task_id=654,
    )

    assert same_workspace is workspace
    assert same_workspace.db_task_id == 654
    assert same_workspace.current_task_id == 654


@pytest.mark.parametrize(
    ("durable_segments", "expected"),
    [
        (None, ("tenant-a",)),
        ((), ()),
    ],
)
def test_agent_service_workspace_preserves_config_db_task_id(tmp_path, durable_segments, expected):
    class ToolConfig:
        _workspace_config = {
            "base_dir": str(tmp_path / "workspaces"),
            "task_id": "agent_2_abcd1234",
            "db_task_id": 987,
            "user_id": 7,
            "scope_segments": ("tenant-a",),
            "__xagent_file_operation_access_version": 1,
            **({} if durable_segments is None else {"durable_storage_segments": durable_segments}),
        }

        def get_allowed_skills(self):
            return None

    service = AgentService(
        name="worker",
        id="agent_2_abcd1234",
        tool_config=ToolConfig(),
        enable_workspace=True,
    )

    assert service.workspace is not None
    assert service.workspace.id == "agent_2_abcd1234"
    assert service.workspace.db_task_id == 987
    assert service.workspace.current_task_id == 987
    assert service.workspace.owner_user_id == 7
    assert service.workspace.file_operation_access_version == 1
    assert service.workspace.durable_storage_segments == expected


def test_workspace_register_file_stages_then_uses_already_durable_upsert(
    monkeypatch, tmp_path, mock_workspace_db
):
    # Override the global autouse fixture from tests/conftest.py for this module.
    del mock_workspace_db
    object_root = tmp_path / "objects"
    monkeypatch.setenv("XAGENT_FILE_STORAGE_URI", object_root.as_uri())
    get_unscoped_file_storage.cache_clear()

    stage_calls = []
    upsert_calls = []

    from xagent.web.services import uploaded_file_store as store_module

    original_stage = store_module.stage_uploaded_file_from_local_path
    original_upsert = store_module.UploadedFileStore.upsert_already_durable

    def stage_spy(**kwargs):
        stage_calls.append(kwargs)
        return original_stage(**kwargs)

    def upsert_spy(self, staged, *, expected, allow_task_rebind=False):
        upsert_calls.append((staged, expected, allow_task_rebind))
        return original_upsert(
            self,
            staged,
            expected=expected,
            allow_task_rebind=allow_task_rebind,
        )

    monkeypatch.setattr(store_module, "stage_uploaded_file_from_local_path", stage_spy)
    monkeypatch.setattr(
        store_module.UploadedFileStore,
        "upsert_already_durable",
        upsert_spy,
    )

    engine = create_engine("sqlite:///:memory:", connect_args={"check_same_thread": False})
    SessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)
    Base.metadata.create_all(bind=engine)
    db = SessionLocal()
    try:
        user = User(username="workspace-user", password_hash="hash")
        db.add(user)
        db.flush()
        task = Task(id=456, user_id=user.id, title="Workspace task")
        db.add(task)
        db.commit()

        workspace = TaskWorkspace(id="web_task_456", base_dir=str(tmp_path / "workspaces"))
        output_path = workspace.output_dir / "report.txt"
        output_path.write_text("workspace output", encoding="utf-8")

        file_id = workspace.register_file(str(output_path), db_session=db)
        db.commit()

        assert len(stage_calls) == 1
        assert stage_calls[0]["local_path"] == output_path
        assert stage_calls[0]["user_id"] == user.id
        assert stage_calls[0]["file_id"] == file_id
        assert stage_calls[0]["task_id"] == 456
        assert stage_calls[0]["workspace_relative_path"] == "output/report.txt"
        assert stage_calls[0]["workspace_category"] == "output"
        _assert_task_output_generation_key(
            stage_calls[0]["storage_key"],
            user_id=int(user.id),
            task_id=456,
            file_id=file_id,
            relative_path="output/report.txt",
        )
        assert len(upsert_calls) == 1
        staged, expected, allow_task_rebind = upsert_calls[0]
        assert staged.storage_key == stage_calls[0]["storage_key"]
        assert expected is None
        assert allow_task_rebind is False
    finally:
        db.close()
        engine.dispose()


def test_workspace_register_file_resyncs_existing_modified_file(
    monkeypatch, tmp_path, mock_workspace_db
):
    # Override the global autouse fixture from tests/conftest.py for this module.
    del mock_workspace_db
    object_root = tmp_path / "objects"
    monkeypatch.setenv("XAGENT_FILE_STORAGE_URI", object_root.as_uri())
    get_unscoped_file_storage.cache_clear()

    engine = create_engine("sqlite:///:memory:", connect_args={"check_same_thread": False})
    SessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)
    Base.metadata.create_all(bind=engine)
    db = SessionLocal()
    try:
        user = User(username="workspace-user", password_hash="hash")
        db.add(user)
        db.flush()
        task = Task(id=654, user_id=user.id, title="Workspace task")
        db.add(task)
        db.commit()

        workspace = TaskWorkspace(id="web_task_654", base_dir=str(tmp_path / "workspaces"))
        output_path = workspace.output_dir / "report.txt"
        output_path.write_text("old", encoding="utf-8")

        file_id = workspace.register_file(str(output_path), db_session=db)
        db.commit()

        output_path.write_text("new content", encoding="utf-8")
        second_file_id = workspace.register_file(str(output_path), db_session=db)
        db.commit()

        assert second_file_id == file_id
        record = db.query(UploadedFile).filter(UploadedFile.file_id == file_id).one()
        assert record.file_size == len("new content")
        assert record.storage_status == "available"

        assert (object_root / str(record.storage_key)).read_text(encoding="utf-8") == "new content"
    finally:
        db.close()
        engine.dispose()


def test_auto_register_files_resyncs_modified_existing_file(
    monkeypatch, tmp_path, mock_workspace_db
):
    # Override the global autouse fixture from tests/conftest.py for this module.
    del mock_workspace_db
    object_root = tmp_path / "objects"
    monkeypatch.setenv("XAGENT_FILE_STORAGE_URI", object_root.as_uri())
    get_unscoped_file_storage.cache_clear()

    engine = create_engine("sqlite:///:memory:", connect_args={"check_same_thread": False})
    SessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)
    Base.metadata.create_all(bind=engine)
    db = SessionLocal()
    try:
        user = User(username="workspace-user", password_hash="hash")
        db.add(user)
        db.flush()
        task = Task(id=655, user_id=user.id, title="Workspace task")
        db.add(task)
        db.commit()

        workspace = TaskWorkspace(id="web_task_655", base_dir=str(tmp_path / "workspaces"))
        output_path = workspace.output_dir / "report.txt"
        output_path.write_text("old", encoding="utf-8")
        file_id = workspace.register_file(str(output_path), db_session=db)
        db.commit()

        workspace.db_session = db
        with workspace.auto_register_files():
            output_path.write_text("new content", encoding="utf-8")
        db.commit()

        record = db.query(UploadedFile).filter(UploadedFile.file_id == file_id).one()
        assert record.file_size == len("new content")
        assert (object_root / str(record.storage_key)).read_text(encoding="utf-8") == "new content"
    finally:
        db.close()
        engine.dispose()


def test_auto_register_files_isolates_one_file_registration_failure(
    monkeypatch,
    tmp_path,
    mock_workspace_db,
    constrained_workspace_db,
):
    del mock_workspace_db
    _engine, _SessionLocal, db = constrained_workspace_db
    _seed_workspace_task(
        db,
        task_id=9010,
        username="workspace-auto-register-isolation-user",
    )
    object_root = tmp_path / "objects"
    monkeypatch.setenv("XAGENT_FILE_STORAGE_URI", object_root.as_uri())
    get_unscoped_file_storage.cache_clear()

    workspace = TaskWorkspace(
        id="web_task_9010",
        base_dir=str(tmp_path / "workspaces"),
    )
    workspace.db_session = db
    original_describe = workspace.describe_file_registration

    def describe_with_one_failure(file_path: str):
        if file_path.endswith("bad.txt"):
            raise RuntimeError("injected registration failure")
        return original_describe(file_path)

    monkeypatch.setattr(
        workspace,
        "describe_file_registration",
        describe_with_one_failure,
    )

    with workspace.auto_register_files():
        (workspace.output_dir / "bad.txt").write_text("bad", encoding="utf-8")
        (workspace.output_dir / "good.txt").write_text("good", encoding="utf-8")

    records = db.query(UploadedFile).all()
    assert [record.filename for record in records] == ["good.txt"]


def test_register_files_deduplicates_one_canonical_path(
    monkeypatch,
    tmp_path,
    mock_workspace_db,
    constrained_workspace_db,
):
    del mock_workspace_db
    _engine, _SessionLocal, db = constrained_workspace_db
    _seed_workspace_task(
        db,
        task_id=9011,
        username="workspace-batch-dedup-user",
    )
    object_root = tmp_path / "objects"
    monkeypatch.setenv("XAGENT_FILE_STORAGE_URI", object_root.as_uri())
    get_unscoped_file_storage.cache_clear()

    workspace = TaskWorkspace(
        id="web_task_9011",
        base_dir=str(tmp_path / "workspaces"),
    )
    output_path = workspace.output_dir / "report.txt"
    output_path.write_text("one artifact", encoding="utf-8")

    file_ids = workspace.register_files(
        (
            (str(output_path), None),
            (str(output_path.resolve()), None),
        ),
        db_session=db,
    )

    assert len(file_ids) == 2
    assert file_ids[0] == file_ids[1]
    assert db.query(UploadedFile).count() == 1


@pytest.mark.parametrize("conflict", ["path", "file_id"])
def test_register_files_rejects_conflicting_batch_identity(
    tmp_path,
    conflict: str,
) -> None:
    workspace = TaskWorkspace(
        id="non_db_workspace",
        base_dir=str(tmp_path / "workspaces"),
    )
    first_path = workspace.output_dir / "first.txt"
    second_path = workspace.output_dir / "second.txt"
    first_path.write_text("first", encoding="utf-8")
    second_path.write_text("second", encoding="utf-8")
    files = (
        ((str(first_path), "file-a"), (str(first_path), "file-b"))
        if conflict == "path"
        else ((str(first_path), "file-a"), (str(second_path), "file-a"))
    )

    with pytest.raises(ValueError):
        workspace.register_files(files)


def test_concurrent_registration_of_one_path_uses_one_workspace_owner(
    monkeypatch,
    tmp_path,
    mock_workspace_db,
):
    del mock_workspace_db
    engine = create_engine(
        f"sqlite:///{tmp_path / 'workspace-concurrency.db'}",
        connect_args={"check_same_thread": False},
        poolclass=QueuePool,
        pool_size=4,
        max_overflow=0,
    )
    SessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)
    Base.metadata.create_all(bind=engine)
    with SessionLocal() as db:
        _seed_workspace_task(
            db,
            task_id=9012,
            username="workspace-concurrent-registration-user",
        )

    from xagent.web.models import database as database_module

    monkeypatch.setattr(database_module, "get_session_local", lambda: SessionLocal)
    object_root = tmp_path / "objects"
    monkeypatch.setenv("XAGENT_FILE_STORAGE_URI", object_root.as_uri())
    get_unscoped_file_storage.cache_clear()

    workspace = TaskWorkspace(
        id="web_task_9012",
        base_dir=str(tmp_path / "workspaces"),
    )
    output_path = workspace.output_dir / "report.txt"
    output_path.write_text("one artifact", encoding="utf-8")
    first_entered = Event()
    second_attempted = Event()
    release_first = Event()
    attempts_lock = Lock()
    attempts = 0
    registration_lock = workspace._registration_lock
    ownership = local()
    original_load = workspace._load_file_registration_plans

    def checked_load(*args, **kwargs):
        assert getattr(ownership, "depth", 0) > 0, "registration read must hold the lock"
        return original_load(*args, **kwargs)

    monkeypatch.setattr(workspace, "_load_file_registration_plans", checked_load)

    class ObservedRegistrationLock:
        def __enter__(self):
            nonlocal attempts
            with attempts_lock:
                attempts += 1
                attempt = attempts
            if attempt == 2:
                acquired = registration_lock.acquire(blocking=False)
                if acquired:
                    registration_lock.release()
                assert not acquired, "the first registration must still hold the lock"
                second_attempted.set()
            registration_lock.acquire()
            if attempt == 1:
                first_entered.set()
                if not release_first.wait(timeout=30):
                    registration_lock.release()
                    raise AssertionError("first registration was never released")
            ownership.depth = getattr(ownership, "depth", 0) + 1
            return self

        def __exit__(self, *_exc):
            ownership.depth -= 1
            registration_lock.release()

    monkeypatch.setattr(workspace, "_registration_lock", ObservedRegistrationLock())
    try:
        with ThreadPoolExecutor(max_workers=2) as executor:
            first = executor.submit(workspace.register_file, str(output_path))
            try:
                assert first_entered.wait(timeout=30)
                second = executor.submit(workspace.register_file, str(output_path))
                assert second_attempted.wait(timeout=30)
                assert not first.done()
                assert not second.done()
            finally:
                release_first.set()
            file_ids = [first.result(timeout=30), second.result(timeout=30)]

        assert file_ids[0] == file_ids[1]
        with SessionLocal() as verify_db:
            assert verify_db.query(UploadedFile).count() == 1
    finally:
        engine.dispose()


def test_auto_register_files_does_not_suppress_body_exception(tmp_path):
    workspace = TaskWorkspace(
        id="non_db_workspace",
        base_dir=str(tmp_path / "workspaces"),
    )

    with pytest.raises(RuntimeError, match="tool execution failed"):
        with workspace.auto_register_files():
            raise RuntimeError("tool execution failed")


def test_workspace_register_file_resyncs_external_file_without_reclassifying_upload(
    monkeypatch, tmp_path, mock_workspace_db
):
    # Override the global autouse fixture from tests/conftest.py for this module.
    del mock_workspace_db
    object_root = tmp_path / "objects"
    monkeypatch.setenv("XAGENT_FILE_STORAGE_URI", object_root.as_uri())
    get_unscoped_file_storage.cache_clear()

    engine = create_engine("sqlite:///:memory:", connect_args={"check_same_thread": False})
    SessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)
    Base.metadata.create_all(bind=engine)
    db = SessionLocal()
    try:
        user = User(username="workspace-user", password_hash="hash")
        db.add(user)
        db.flush()
        task = Task(id=656, user_id=user.id, title="Workspace task")
        db.add(task)
        db.commit()

        external_dir = tmp_path / "external-uploads"
        external_dir.mkdir()
        external_path = external_dir / "source.txt"
        external_path.write_text("old upload", encoding="utf-8")

        from xagent.web.services.uploaded_file_store import UploadedFileStore

        record = UploadedFileStore(db).create_from_local_path(
            local_path=external_path,
            user_id=int(user.id),
            task_id=int(task.id),
            filename="source.txt",
            mime_type="text/plain",
        )
        file_id = str(record.file_id)
        original_storage_key = str(record.storage_key)
        db.commit()

        workspace = TaskWorkspace(
            id="web_task_656",
            base_dir=str(tmp_path / "workspaces"),
            allowed_external_dirs=[str(external_dir)],
        )

        external_path.write_text("new upload", encoding="utf-8")
        second_file_id = workspace.register_file(str(external_path), file_id=file_id, db_session=db)
        db.commit()

        assert second_file_id == file_id
        db.refresh(record)
        assert record.storage_key != original_storage_key
        _assert_task_output_generation_key(
            str(record.storage_key),
            user_id=int(user.id),
            task_id=656,
            file_id=file_id,
            relative_path="source.txt",
        )
        assert record.workspace_relative_path is None
        assert record.workspace_category is None
        assert record.file_size == len("new upload")

        assert (object_root / str(record.storage_key)).read_text(encoding="utf-8") == "new upload"
    finally:
        db.close()
        engine.dispose()


def test_list_all_user_files_includes_durable_only_uploads(tmp_path, mock_workspace_db):
    # Override the global autouse fixture from tests/conftest.py for this module.
    del mock_workspace_db

    engine = create_engine("sqlite:///:memory:", connect_args={"check_same_thread": False})
    SessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)
    Base.metadata.create_all(bind=engine)
    db = SessionLocal()
    try:
        user = User(username="workspace-user", password_hash="hash")
        db.add(user)
        db.flush()
        task = Task(id=789, user_id=user.id, title="Workspace task")
        db.add(task)
        db.commit()

        missing_local_path = tmp_path / "uploads" / "durable-only.txt"
        assert not missing_local_path.exists()
        file_record = UploadedFile(
            user_id=user.id,
            task_id=task.id,
            filename="durable-only.txt",
            storage_path=str(missing_local_path),
            storage_backend="s3",
            storage_key=f"users/{user.id}/uploads/file-1/durable-only.txt",
            storage_status="available",
            mime_type="text/plain",
            file_size=12,
        )
        db.add(file_record)
        db.commit()
        db.refresh(file_record)

        workspace = TaskWorkspace(
            id="web_task_789",
            base_dir=str(tmp_path / "workspaces"),
        )
        workspace.db_session = db

        result = workspace.list_all_user_files(include_workspace_files=False)

        assert result["success"] is True
        assert [file_info["file_id"] for file_info in result["files"]] == [file_record.file_id]
        assert result["files"][0]["filename"] == "durable-only.txt"
        assert result["files"][0]["in_current_workspace"] is False
    finally:
        db.close()
        engine.dispose()


@pytest.fixture
def mock_workspace_db():
    yield


def test_stage_and_discard_external_upload_file(tmp_path, monkeypatch):
    workspace = TaskWorkspace(id="task_stage", base_dir=str(tmp_path / "workspaces"))
    source = tmp_path / "materialized" / "report.xlsx"
    source.parent.mkdir()
    source.write_bytes(b"workbook")
    monkeypatch.setattr(workspace, "resolve_file_id_detached", lambda _: source)

    staged = workspace.stage_file_for_external_upload("file-id")

    assert staged.is_file()
    assert staged.read_bytes() == b"workbook"
    assert staged.is_relative_to(workspace.temp_dir / ".xagent-internal" / "mcp-upload")
    assert str(staged) not in workspace.get_allowed_dirs()
    assert any(staged.is_relative_to(Path(directory)) for directory in workspace.get_allowed_dirs())

    workspace.discard_staged_external_upload(staged)
    assert not staged.exists()
