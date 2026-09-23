import io
import json
import zipfile
from unittest.mock import MagicMock, Mock

import pytest
import requests
from lxml import etree
from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from xagent.core.tools.adapters.vibe.output_filter import OutputValueFilter
from xagent.web.tools.mcp import powerpoint


def _pptx_bytes(build_fn=None) -> bytes:
    presentation = Presentation()
    if build_fn:
        build_fn(presentation)
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


class MockResponse:
    def __init__(self, json_data=None, status_code=200, content=None, url=None):
        self._json_data = json_data if json_data is not None else {}
        self.status_code = status_code
        self.content = (
            json.dumps(self._json_data).encode("utf-8") if content is None else content
        )
        self.text = self.content.decode("utf-8", errors="replace")
        self.url = url or "https://graph.microsoft.com/v1.0/example"

    def json(self):
        return self._json_data

    def raise_for_status(self):
        if self.status_code >= 400:
            raise requests.HTTPError(
                f"{self.status_code} Client Error: Error for url: {self.url}",
                response=self,
            )

    def iter_content(self, chunk_size=1):
        for start in range(0, len(self.content), chunk_size):
            yield self.content[start : start + chunk_size]

    def close(self):
        return None


def _metadata_response(content: bytes, etag: str = '"etag-1"') -> MockResponse:
    return MockResponse(
        {
            "id": "item-1",
            "size": len(content),
            "eTag": etag,
            "@microsoft.graph.downloadUrl": "https://download.example/deck.pptx",
        }
    )


def _mock_download(monkeypatch, content: bytes, etag: str = '"etag-1"'):
    mock_request = Mock(return_value=_metadata_response(content, etag))
    mock_get = Mock(return_value=MockResponse(content=content))
    monkeypatch.setattr(powerpoint.requests, "request", mock_request)
    monkeypatch.setattr(powerpoint.requests, "get", mock_get)
    return mock_request, mock_get


def _mock_versioned_write(monkeypatch, content: bytes, etag: str = '"etag-1"'):
    responses = iter(
        [
            _metadata_response(content, etag),
            MockResponse({"uploadUrl": "https://upload.example/session"}),
        ]
    )
    mock_request = Mock(side_effect=lambda *a, **k: next(responses))
    mock_get = Mock(return_value=MockResponse(content=content))
    mock_put = Mock(return_value=MockResponse({"id": "item-1"}, status_code=200))
    mock_session_cls = MagicMock()
    mock_session_cls.return_value.__enter__.return_value.put = mock_put
    monkeypatch.setattr(powerpoint.requests, "request", mock_request)
    monkeypatch.setattr(powerpoint.requests, "get", mock_get)
    monkeypatch.setattr(powerpoint.requests, "Session", mock_session_cls)
    return mock_request, mock_get, mock_put


@pytest.fixture(autouse=True)
def _credentials(monkeypatch):
    monkeypatch.setenv("AUTH_TOKEN", "test-graph-token")


@pytest.mark.parametrize(
    ("response_factory", "expected_status"),
    [
        (
            lambda: powerpoint._success(
                item={"id": "x" * 500, "description": "y" * 500},
                slide_index=3,
            ),
            "success",
        ),
        (lambda: powerpoint._error("x" * 500, details={"raw": "y" * 500}), "error"),
        (lambda: powerpoint._conflict("x" * 500), "conflict"),
        (lambda: powerpoint._indeterminate("x" * 500), "indeterminate"),
    ],
)
def test_outcome_envelopes_survive_platform_string_filter(
    monkeypatch, response_factory, expected_status
):
    """Mutation outcomes must be valid JSON after the generic output filter.

    That filter truncates strings without understanding JSON, so PowerPoint
    must make each envelope fit before returning it.
    """
    max_chars = 80
    monkeypatch.setattr(powerpoint, "get_tool_max_output_length", lambda: max_chars)

    raw = response_factory()
    filtered = OutputValueFilter(
        max_chars=max_chars, max_fields=100, max_recursion=10
    ).filter(raw, "powerpoint_mutation")

    assert filtered == raw
    assert len(raw) <= max_chars
    result = json.loads(filtered)
    assert result["status"] == expected_status
    if expected_status == "indeterminate":
        assert result["safe_to_retry"] is False


def test_success_envelope_preserves_provider_metadata_when_it_fits(monkeypatch):
    monkeypatch.setattr(powerpoint, "get_tool_max_output_length", lambda: 500)
    item = {"id": "item-1", "customFacet": {"value": "kept"}}

    result = json.loads(powerpoint._success(item=item))

    assert result["item"] == item


# ---------------------------------------------------------------------------
# path helpers
# ---------------------------------------------------------------------------


def test_item_path_defaults_to_own_onedrive():
    assert (
        powerpoint._item_path("Deck.pptx", None, None) == "/me/drive/root:/Deck.pptx:"
    )


def test_item_path_sharepoint_path_form_site_id_closes_colon_before_drive():
    """A "hostname:/server-relative-path" site_id must get a second,
    closing colon before /drive -- Graph's documented example is
    ".../sites/contoso.sharepoint.com:/teams/hr:/drive"; without it Graph
    parses "/drive" as part of the site-relative path instead."""
    assert (
        powerpoint._item_path("Deck.pptx", "contoso.sharepoint.com:/teams/hr", None)
        == "/sites/contoso.sharepoint.com:/teams/hr:/drive/root:/Deck.pptx:"
    )


def test_item_path_sharepoint_path_form_site_id_with_drive_id():
    assert (
        powerpoint._item_path(
            "Deck.pptx", "contoso.sharepoint.com:/teams/hr", "drive-1"
        )
        == "/sites/contoso.sharepoint.com:/teams/hr:/drives/drive-1/root:/Deck.pptx:"
    )


def test_item_path_sharepoint_composite_site_id_unaffected():
    """The composite "hostname,spSiteId,spWebId" form has no ':' and must
    not get an extra colon inserted."""
    assert (
        powerpoint._item_path("Deck.pptx", "contoso.sharepoint.com,site,web", None)
        == "/sites/contoso.sharepoint.com,site,web/drive/root:/Deck.pptx:"
    )


def test_content_path_appends_content():
    assert (
        powerpoint._content_path("Deck.pptx", None, None)
        == "/me/drive/root:/Deck.pptx:/content"
    )


def test_normalize_relative_path_rejects_trailing_period():
    with pytest.raises(ValueError, match="must not end with a period"):
        powerpoint._normalize_relative_path("Deck.pptx.")


@pytest.mark.parametrize("path", ["Deck.pptm", "Deck.ppt", "Deck", "Deck.ppsx"])
def test_normalize_relative_path_rejects_unsupported_powerpoint_package(path):
    with pytest.raises(ValueError, match=r"must name a \.pptx"):
        powerpoint._normalize_relative_path(path)


def test_normalize_relative_path_allows_case_insensitive_pptx_extension():
    assert powerpoint._normalize_relative_path("Deck.PPTX") == "Deck.PPTX"


def test_normalize_relative_path_rejects_dot_segments():
    with pytest.raises(ValueError, match="must not contain"):
        powerpoint._normalize_relative_path("../secret.pptx")


def test_normalize_relative_path_rejects_doubled_slash():
    """An empty segment (from a doubled '/') isn't '.' or '..' but would
    still build a malformed root:/{path}: Graph URL if let through."""
    with pytest.raises(ValueError, match="empty"):
        powerpoint._normalize_relative_path("reports//Q1.pptx")


def test_site_segment_rejects_doubled_slash():
    with pytest.raises(ValueError, match="empty"):
        powerpoint._site_segment("contoso.sharepoint.com:/teams//hr")


def test_item_path_rejects_empty_site_id():
    """An empty string is a caller mistake (e.g. an upstream field
    defaulting unset to "" instead of None), not "not provided" -- it must
    not be silently treated the same as site_id=None and routed to
    /me/drive instead."""
    with pytest.raises(ValueError, match="site_id is required"):
        powerpoint._item_path("Deck.pptx", "", None)


def test_item_path_rejects_empty_drive_id():
    with pytest.raises(ValueError, match="drive_id"):
        powerpoint._item_path("Deck.pptx", None, "")


# ---------------------------------------------------------------------------
# slide helpers
# ---------------------------------------------------------------------------


def test_shape_text_returns_none_for_non_text_shape():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    picture_placeholder_shapes = list(slide.shapes)
    # A blank layout has no shapes at all; simulate a shape without a text
    # frame via a connector, the simplest shape type with has_text_frame=False.
    from pptx.enum.shapes import MSO_CONNECTOR

    connector = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(0), Inches(0), Inches(1), Inches(1)
    )
    assert not connector.has_text_frame
    assert powerpoint._shape_text(connector) is None
    assert picture_placeholder_shapes == []


def test_require_slide_out_of_range():
    presentation = Presentation()
    with pytest.raises(ValueError, match="out of range"):
        powerpoint._require_slide(presentation, 0)


def test_require_int_accepts_int():
    assert powerpoint._require_int(3, "slide_index") == 3


def test_require_int_rejects_non_int():
    with pytest.raises(TypeError, match="slide_index must be an integer"):
        powerpoint._require_int("0", "slide_index")


def test_require_int_rejects_bool():
    """bool is a subclass of int in Python -- isinstance(True, int) is True
    -- so this must be checked explicitly rather than relying on isinstance
    alone, since True/False are never valid slide/shape/layout indices."""
    with pytest.raises(TypeError, match="slide_index must be an integer"):
        powerpoint._require_int(True, "slide_index")


def _make_placeholder(idx: int, has_text_frame: bool) -> Mock:
    placeholder = Mock()
    placeholder.placeholder_format.idx = idx
    placeholder.has_text_frame = has_text_frame
    return placeholder


def test_select_body_placeholder_skips_title():
    title = _make_placeholder(0, has_text_frame=True)
    body = _make_placeholder(1, has_text_frame=True)
    slide = Mock()
    slide.placeholders = [title, body]
    assert powerpoint._select_body_placeholder(slide) is body


def test_select_body_placeholder_skips_non_text_placeholder():
    """A picture/chart placeholder can sit at a lower idx than the real
    text placeholder -- this must not stop at it and give up."""
    title = _make_placeholder(0, has_text_frame=True)
    picture = _make_placeholder(1, has_text_frame=False)
    body = _make_placeholder(2, has_text_frame=True)
    slide = Mock()
    slide.placeholders = [title, picture, body]
    assert powerpoint._select_body_placeholder(slide) is body


def test_select_body_placeholder_returns_none_when_no_match():
    title = _make_placeholder(0, has_text_frame=True)
    picture = _make_placeholder(1, has_text_frame=False)
    slide = Mock()
    slide.placeholders = [title, picture]
    assert powerpoint._select_body_placeholder(slide) is None


def test_delete_slide_removes_correct_slide():
    presentation = Presentation()
    for title in ("A", "B", "C"):
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = title

    powerpoint._delete_slide(presentation, 1)

    assert [s.shapes.title.text for s in presentation.slides] == ["A", "C"]


def test_delete_slide_out_of_range():
    presentation = Presentation()
    with pytest.raises(ValueError, match="out of range"):
        powerpoint._delete_slide(presentation, 0)


def test_delete_slide_prunes_relationship_and_part():
    """Removing only the <p:sldId> entry leaves the slide's own part and its
    relationship in the saved package; _delete_slide must also drop the
    relationship so the "deleted" slide's content doesn't survive a
    save/reload round trip at all."""
    presentation = Presentation()
    for title in ("A", "B"):
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = title
    rels_before = set(presentation.part.rels)

    powerpoint._delete_slide(presentation, 0)

    rels_after = set(presentation.part.rels)
    assert rels_after < rels_before
    assert len(rels_before) - len(rels_after) == 1

    buffer = io.BytesIO()
    presentation.save(buffer)
    reloaded = Presentation(io.BytesIO(buffer.getvalue()))
    assert [s.shapes.title.text for s in reloaded.slides] == ["B"]


def test_delete_slide_rejects_custom_show_reference_without_mutating():
    presentation = Presentation()
    target_slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    target = list(presentation.slides._sldIdLst)[0]
    relationship_id = target.get(qn("r:id"))
    custom_show_list = etree.SubElement(presentation._element, qn("p:custShowLst"))
    custom_show = etree.SubElement(custom_show_list, qn("p:custShow"))
    custom_show.set("name", "Referenced slide")
    custom_show.set("id", "1")
    custom_slide_list = etree.SubElement(custom_show, qn("p:sldLst"))
    custom_slide = etree.SubElement(custom_slide_list, qn("p:sld"))
    custom_slide.set(qn("r:id"), relationship_id)

    with pytest.raises(ValueError, match="custom slide show"):
        powerpoint._delete_slide(presentation, 0)

    assert list(presentation.slides) == [target_slide]
    assert relationship_id in presentation.part.rels


def test_delete_slide_rejects_link_from_another_slide_without_mutating():
    presentation = Presentation()
    source = presentation.slides.add_slide(presentation.slide_layouts[6])
    target = presentation.slides.add_slide(presentation.slide_layouts[6])
    source.part.relate_to(target.part, RT.SLIDE)

    with pytest.raises(ValueError, match="linked from another slide"):
        powerpoint._delete_slide(presentation, 1)

    assert list(presentation.slides) == [source, target]


def test_delete_slide_rejects_section_reference_without_mutating():
    presentation = Presentation()
    target_slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    target = list(presentation.slides._sldIdLst)[0]
    powerpoint_2010 = "http://schemas.microsoft.com/office/powerpoint/2010/main"
    extension_list = etree.SubElement(presentation._element, qn("p:extLst"))
    extension = etree.SubElement(extension_list, qn("p:ext"))
    extension.set("uri", "{521415D9-36F7-43E2-AB2F-B90AF26B5E84}")
    section_list = etree.SubElement(extension, f"{{{powerpoint_2010}}}sectionLst")
    section = etree.SubElement(section_list, f"{{{powerpoint_2010}}}section")
    section.set("name", "Section 1")
    section.set("id", "{00000000-0000-0000-0000-000000000001}")
    section_slide_list = etree.SubElement(section, f"{{{powerpoint_2010}}}sldIdLst")
    section_slide_id = etree.SubElement(
        section_slide_list, f"{{{powerpoint_2010}}}sldId"
    )
    section_slide_id.set("id", target.get("id"))

    with pytest.raises(ValueError, match="presentation section"):
        powerpoint._delete_slide(presentation, 0)

    assert list(presentation.slides) == [target_slide]


def test_delete_slide_with_notes_remains_supported():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.notes_slide.notes_text_frame.text = "Speaker notes"

    powerpoint._delete_slide(presentation, 0)

    assert len(presentation.slides) == 0
    buffer = io.BytesIO()
    presentation.save(buffer)
    assert len(Presentation(io.BytesIO(buffer.getvalue())).slides) == 0


# ---------------------------------------------------------------------------
# create
# ---------------------------------------------------------------------------


def test_create_presentation_rejects_unsupported_extension_before_graph(monkeypatch):
    mock_request = Mock()
    monkeypatch.setattr(powerpoint.requests, "request", mock_request)

    result = json.loads(powerpoint.powerpoint_create_presentation("Deck.pptm"))

    assert result["status"] == "error"
    assert ".pptx" in result["message"]
    mock_request.assert_not_called()


def test_create_presentation_rejects_when_session_creation_conflicts(monkeypatch):
    mock_request = Mock(
        return_value=MockResponse(
            {"error": {"code": "nameAlreadyExists"}}, status_code=409
        )
    )
    monkeypatch.setattr(powerpoint.requests, "request", mock_request)

    result = json.loads(powerpoint.powerpoint_create_presentation("Deck.pptx"))

    assert result["status"] == "error"
    assert "already exists" in result["message"]
    assert mock_request.call_args.kwargs["url"].endswith("createUploadSession")


def test_create_presentation_keeps_final_upload_conflict_definite(monkeypatch):
    mock_request = Mock(
        return_value=MockResponse({"uploadUrl": "https://upload.example/session"})
    )
    monkeypatch.setattr(powerpoint.requests, "request", mock_request)
    monkeypatch.setattr(
        powerpoint.requests,
        "put",
        Mock(return_value=MockResponse({}, status_code=409)),
    )
    # Cleanup is best-effort and must not replace the original definite error,
    # even if the provider rejects or loses the DELETE too.
    mock_delete = Mock(side_effect=requests.ConnectionError("cleanup failed"))
    monkeypatch.setattr(powerpoint.requests, "delete", mock_delete)

    result = json.loads(powerpoint.powerpoint_create_presentation("Deck.pptx"))

    assert result["status"] == "error"
    assert "already exists" in result["message"]
    mock_delete.assert_called_once_with(
        "https://upload.example/session",
        timeout=powerpoint._UPLOAD_CANCEL_TIMEOUT_SECONDS,
    )


def test_create_only_upload_cancels_session_after_definite_http_failure(monkeypatch):
    secret_url = "https://upload.example/session?token=super-secret-value"
    monkeypatch.setattr(
        powerpoint.requests,
        "request",
        Mock(return_value=MockResponse({"uploadUrl": secret_url})),
    )
    monkeypatch.setattr(
        powerpoint.requests,
        "put",
        Mock(return_value=MockResponse({}, status_code=400, url=secret_url)),
    )
    mock_delete = Mock(return_value=MockResponse({}, status_code=204))
    monkeypatch.setattr(powerpoint.requests, "delete", mock_delete)

    with pytest.raises(powerpoint._GraphRequestError) as exc_info:
        powerpoint._create_only_upload(b"content", "Deck.pptx", None, None)

    assert "super-secret-value" not in str(exc_info.value)
    mock_delete.assert_called_once_with(
        secret_url, timeout=powerpoint._UPLOAD_CANCEL_TIMEOUT_SECONDS
    )


def test_create_presentation_upload_failure_does_not_leak_session_url(monkeypatch):
    """The upload-session URL is a pre-authenticated bearer secret (a token
    in its own query string) -- a failed upload must never let that URL
    reach the caller through the error message, whether the failure comes
    from raise_for_status() or a lower-level connection/timeout error."""
    secret_url = "https://upload.example/session?token=super-secret-value"
    mock_request = Mock(return_value=MockResponse({"uploadUrl": secret_url}))
    monkeypatch.setattr(powerpoint.requests, "request", mock_request)
    mock_put = Mock(return_value=MockResponse({}, status_code=500, url=secret_url))
    monkeypatch.setattr(powerpoint.requests, "put", mock_put)

    result = json.loads(powerpoint.powerpoint_create_presentation("Deck.pptx"))

    assert result["status"] == "indeterminate"
    assert result["safe_to_retry"] is False
    assert "super-secret-value" not in result["message"]
    assert secret_url not in result["message"]


def test_create_presentation_upload_connection_error_does_not_leak_session_url(
    monkeypatch,
):
    secret_url = "https://upload.example/session?token=super-secret-value"
    mock_request = Mock(return_value=MockResponse({"uploadUrl": secret_url}))
    monkeypatch.setattr(powerpoint.requests, "request", mock_request)
    mock_put = Mock(
        side_effect=requests.ConnectionError(f"Connection refused: {secret_url}")
    )
    monkeypatch.setattr(powerpoint.requests, "put", mock_put)

    result = json.loads(powerpoint.powerpoint_create_presentation("Deck.pptx"))

    assert result["status"] == "indeterminate"
    assert result["safe_to_retry"] is False
    assert "super-secret-value" not in result["message"]
    assert secret_url not in result["message"]


def test_create_only_upload_does_not_chain_secret_bearing_exception(monkeypatch):
    """raise ... from exc would still attach the original, URL-bearing
    exception as __cause__ even though the raised message itself is
    sanitized -- a future traceback/log/APM capture could surface it.
    from None must be used instead, matching onedrive.py's precedent."""
    secret_url = "https://upload.example/session?token=super-secret-value"
    mock_request = Mock(return_value=MockResponse({"uploadUrl": secret_url}))
    monkeypatch.setattr(powerpoint.requests, "request", mock_request)
    mock_put = Mock(return_value=MockResponse({}, status_code=500, url=secret_url))
    monkeypatch.setattr(powerpoint.requests, "put", mock_put)

    with pytest.raises(powerpoint._IndeterminateWriteError) as exc_info:
        powerpoint._create_only_upload(b"content", "Deck.pptx", None, None)

    assert exc_info.value.__cause__ is None


def test_create_only_upload_connection_error_does_not_chain_exception(monkeypatch):
    secret_url = "https://upload.example/session?token=super-secret-value"
    mock_request = Mock(return_value=MockResponse({"uploadUrl": secret_url}))
    monkeypatch.setattr(powerpoint.requests, "request", mock_request)
    mock_put = Mock(
        side_effect=requests.ConnectionError(f"Connection refused: {secret_url}")
    )
    monkeypatch.setattr(powerpoint.requests, "put", mock_put)

    with pytest.raises(powerpoint._IndeterminateWriteError) as exc_info:
        powerpoint._create_only_upload(b"content", "Deck.pptx", None, None)

    assert exc_info.value.__cause__ is None


def test_create_presentation_uploads_blank_presentation(monkeypatch):
    mock_request = Mock(
        return_value=MockResponse({"uploadUrl": "https://upload.example/session"})
    )
    monkeypatch.setattr(powerpoint.requests, "request", mock_request)
    mock_put = Mock(return_value=MockResponse({"id": "new-item"}, status_code=201))
    monkeypatch.setattr(powerpoint.requests, "put", mock_put)

    result = json.loads(powerpoint.powerpoint_create_presentation("Deck.pptx"))

    assert result["status"] == "success"
    assert result["item"]["id"] == "new-item"
    session_call = mock_request.call_args
    assert session_call.kwargs["json"] == {
        "item": {"@microsoft.graph.conflictBehavior": "fail"}
    }
    put_call = mock_put.call_args
    assert put_call.args[0] == "https://upload.example/session"
    assert (
        put_call.kwargs["headers"]["Content-Type"] == powerpoint._POWERPOINT_MIME_TYPE
    )
    assert "Authorization" not in put_call.kwargs.get("headers", {})
    # Re-parses as a real, valid (if blank) pptx.
    Presentation(io.BytesIO(put_call.kwargs["data"]))


# ---------------------------------------------------------------------------
# read
# ---------------------------------------------------------------------------


def test_get_presentation_text(monkeypatch):
    def build(prs):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "Title A"
        slide.placeholders[1].text_frame.text = "Body A"

    content = _pptx_bytes(build)
    mock_request, mock_get = _mock_download(monkeypatch, content)

    result = json.loads(powerpoint.powerpoint_get_presentation_text("Deck.pptx"))

    assert result["status"] == "success"
    assert result["slides"] == [
        {"slide_index": 0, "shapes": ["Title A", "Body A"], "notes": None}
    ]
    assert result["etag"] == '"etag-1"'
    assert mock_get.call_args.kwargs["timeout"] == powerpoint._BINARY_TIMEOUT_SECONDS
    assert "headers" not in mock_get.call_args.kwargs


def test_get_presentation_text_falls_back_to_authenticated_content_endpoint(
    monkeypatch,
):
    content = _pptx_bytes()
    metadata = _metadata_response(content)
    metadata._json_data.pop("@microsoft.graph.downloadUrl")
    content_response = MockResponse(content=content)
    mock_request = Mock(side_effect=[metadata, content_response])
    monkeypatch.setattr(powerpoint.requests, "request", mock_request)
    mock_get = Mock()
    monkeypatch.setattr(powerpoint.requests, "get", mock_get)

    result = json.loads(powerpoint.powerpoint_get_presentation_text("Deck.pptx"))

    assert result["status"] == "success"
    assert mock_get.call_count == 0
    content_call = mock_request.call_args_list[1]
    assert content_call.kwargs["method"] == "GET"
    assert content_call.kwargs["url"].endswith("/me/drive/root:/Deck.pptx:/content")
    assert content_call.kwargs["headers"]["Authorization"] == "Bearer test-graph-token"
    assert content_call.kwargs["stream"] is True


def test_get_presentation_text_returns_resumable_bounded_pages(monkeypatch):
    def build(prs):
        for index in range(5):
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            slide.shapes.title.text = f"Slide {index} " + "x" * 55

    content = _pptx_bytes(build)
    _mock_download(monkeypatch, content)
    monkeypatch.setattr(powerpoint, "get_tool_max_output_length", lambda: 430)

    cursor = None
    slide_indices = []
    while True:
        result_text = powerpoint.powerpoint_get_presentation_text(
            "Deck.pptx", cursor=cursor
        )
        assert len(result_text) <= 430
        result = json.loads(result_text)
        assert result["status"] == "success"
        slide_indices.extend(slide["slide_index"] for slide in result["slides"])
        if not result["truncated"]:
            assert result["next_cursor"] is None
            break
        cursor = result["next_cursor"]

    assert slide_indices == list(range(5))


def test_get_presentation_text_skips_oversized_slide_and_continues(
    monkeypatch,
):
    def build(prs):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
        box.text_frame.text = "x" * 1000
        later = prs.slides.add_slide(prs.slide_layouts[6])
        later.shapes.add_textbox(
            Inches(1), Inches(1), Inches(3), Inches(1)
        ).text_frame.text = "Later slide"

    content = _pptx_bytes(build)
    _mock_download(monkeypatch, content)
    monkeypatch.setattr(powerpoint, "get_tool_max_output_length", lambda: 430)

    first_text = powerpoint.powerpoint_get_presentation_text("Deck.pptx")
    first = json.loads(first_text)

    assert len(first_text) <= 430
    assert first["status"] == "success"
    assert first["slides"] == []
    assert first["omitted_item"] == {
        "item_index": 0,
        "item_type": "slide",
        "reason": "item_exceeds_output_limit",
        "output_limit": 430,
    }
    assert first["truncated"] is True

    second = json.loads(
        powerpoint.powerpoint_get_presentation_text(
            "Deck.pptx", cursor=first["next_cursor"]
        )
    )
    assert second["slides"] == [
        {"slide_index": 1, "shapes": ["Later slide"], "notes": None}
    ]
    assert second["truncated"] is False


def test_list_slides_cursor_is_rejected_after_presentation_changes(monkeypatch):
    def build(prs):
        for index in range(4):
            prs.slides.add_slide(prs.slide_layouts[1]).shapes.title.text = str(index)

    content = _pptx_bytes(build)
    _mock_download(monkeypatch, content)
    monkeypatch.setattr(powerpoint, "get_tool_max_output_length", lambda: 300)
    first = json.loads(powerpoint.powerpoint_list_slides("Deck.pptx"))
    assert first["truncated"] is True

    _mock_download(monkeypatch, content, etag='"etag-2"')
    second = json.loads(
        powerpoint.powerpoint_list_slides("Deck.pptx", cursor=first["next_cursor"])
    )

    assert second["status"] == "conflict"
    assert "changed while reading" in second["message"]


def test_get_presentation_text_includes_group_table_and_notes(monkeypatch):
    def build(prs):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(1), Inches(1))
        box.text_frame.text = "Top level"
        group = slide.shapes.add_group_shape([box])
        nested_box = group.shapes.add_textbox(
            Inches(2), Inches(2), Inches(1), Inches(1)
        )
        nested_box.text_frame.text = "Nested"
        table_shape = slide.shapes.add_table(
            1, 2, Inches(0), Inches(3), Inches(4), Inches(1)
        )
        table_shape.table.cell(0, 0).text_frame.text = "Cell A"
        table_shape.table.cell(0, 1).text_frame.text = "Cell B"
        slide.notes_slide.notes_text_frame.text = "Speaker notes"

    content = _pptx_bytes(build)
    _mock_download(monkeypatch, content)

    result = json.loads(powerpoint.powerpoint_get_presentation_text("Deck.pptx"))

    assert result["status"] == "success"
    slide_result = result["slides"][0]
    assert "Nested" in slide_result["shapes"]
    assert "Cell A" in slide_result["shapes"]
    assert "Cell B" in slide_result["shapes"]
    assert slide_result["notes"] == "Speaker notes"


def test_slide_notes_returns_none_without_creating_notes_slide():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])

    assert powerpoint._slide_notes(slide) is None
    assert not slide.has_notes_slide


def test_get_presentation_text_rejects_non_pptx_content(monkeypatch):
    _mock_download(monkeypatch, b"not a pptx file")

    result = json.loads(powerpoint.powerpoint_get_presentation_text("Deck.pptx"))

    assert result["status"] == "error"
    assert "PowerPoint presentation" in result["message"]


def test_download_rejects_oversized_metadata_before_fetching_content(monkeypatch):
    metadata = _metadata_response(b"")
    metadata._json_data["size"] = powerpoint._MAX_PRESENTATION_BYTES + 1
    mock_request = Mock(return_value=metadata)
    mock_get = Mock()
    monkeypatch.setattr(powerpoint.requests, "request", mock_request)
    monkeypatch.setattr(powerpoint.requests, "get", mock_get)

    result = json.loads(powerpoint.powerpoint_list_slides("Deck.pptx"))

    assert result["status"] == "error"
    assert "MB limit" in result["message"]
    mock_get.assert_not_called()


def test_download_stream_enforces_byte_limit_when_metadata_is_wrong(monkeypatch):
    monkeypatch.setattr(powerpoint, "_MAX_PRESENTATION_BYTES", 10)
    metadata = _metadata_response(b"small")
    mock_request = Mock(return_value=metadata)
    mock_get = Mock(return_value=MockResponse(content=b"x" * 11))
    monkeypatch.setattr(powerpoint.requests, "request", mock_request)
    monkeypatch.setattr(powerpoint.requests, "get", mock_get)

    result = json.loads(powerpoint.powerpoint_list_slides("Deck.pptx"))

    assert result["status"] == "error"
    assert "download exceeded" in result["message"]


def test_validate_presentation_archive_bounds_expanded_size(monkeypatch):
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("ppt/slides/slide1.xml", b"x" * 100)
    monkeypatch.setattr(powerpoint, "_MAX_PRESENTATION_UNCOMPRESSED_BYTES", 50)

    with pytest.raises(ValueError, match="expands beyond"):
        powerpoint._validate_presentation_archive(buffer.getvalue())


def test_list_slides(monkeypatch):
    def build(prs):
        prs.slides.add_slide(prs.slide_layouts[1])

    content = _pptx_bytes(build)
    _mock_download(monkeypatch, content)

    result = json.loads(powerpoint.powerpoint_list_slides("Deck.pptx"))

    assert result["status"] == "success"
    assert result["slides"][0]["slide_index"] == 0
    assert result["slides"][0]["layout_name"] == "Title and Content"


def test_get_slide_text_skips_oversized_shape_and_continues(monkeypatch):
    def build(prs):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
        box.text_frame.text = "x" * 1000
        slide.shapes.add_textbox(
            Inches(1), Inches(2), Inches(3), Inches(1)
        ).text_frame.text = "Later shape"

    content = _pptx_bytes(build)
    _mock_download(monkeypatch, content)
    monkeypatch.setattr(powerpoint, "get_tool_max_output_length", lambda: 430)

    first_text = powerpoint.powerpoint_get_slide_text("Deck.pptx", 0)
    first = json.loads(first_text)

    assert len(first_text) <= 430
    assert first["status"] == "success"
    assert first["shapes"] == []
    assert first["omitted_item"]["item_index"] == 0
    assert first["omitted_item"]["item_type"] == "shape"
    assert first["truncated"] is True

    second = json.loads(
        powerpoint.powerpoint_get_slide_text(
            "Deck.pptx", 0, cursor=first["next_cursor"]
        )
    )
    assert len(second["shapes"]) == 1
    assert second["shapes"][0]["shape_index"] == 1
    assert second["shapes"][0]["text"] == "Later shape"
    assert second["truncated"] is False


def test_get_slide_text_out_of_range(monkeypatch):
    content = _pptx_bytes()
    _mock_download(monkeypatch, content)

    result = json.loads(powerpoint.powerpoint_get_slide_text("Deck.pptx", 0))

    assert result["status"] == "error"
    assert "out of range" in result["message"]


def test_get_slide_text_rejects_non_int_slide_index(monkeypatch):
    """A direct Python call (e.g. from a test, or any caller bypassing
    FastMCP's own schema validation) with a non-int slide_index must get a
    clean error response, not an unhandled exception -- and must fail before
    ever downloading the presentation."""
    mock_request = Mock()
    monkeypatch.setattr(powerpoint.requests, "request", mock_request)

    result = json.loads(powerpoint.powerpoint_get_slide_text("Deck.pptx", "0"))

    assert result["status"] == "error"
    assert "slide_index must be an integer" in result["message"]
    mock_request.assert_not_called()


def test_list_slide_layouts(monkeypatch):
    content = _pptx_bytes()
    _mock_download(monkeypatch, content)

    result = json.loads(powerpoint.powerpoint_list_slide_layouts("Deck.pptx"))

    assert result["status"] == "success"
    assert result["layouts"][1] == {"layout_index": 1, "name": "Title and Content"}


# ---------------------------------------------------------------------------
# write
# ---------------------------------------------------------------------------


def test_add_slide_sets_title_and_body(monkeypatch):
    content = _pptx_bytes()
    mock_request, _, mock_put = _mock_versioned_write(monkeypatch, content)

    result = json.loads(
        powerpoint.powerpoint_add_slide(
            "Deck.pptx", '"etag-1"', title="New Title", body_text="New Body"
        )
    )

    assert result["status"] == "success"
    assert result["slide_index"] == 0
    session_call = mock_request.call_args_list[1]
    assert session_call.kwargs["headers"]["If-Match"] == '"etag-1"'
    put_call = mock_put.call_args
    uploaded = Presentation(io.BytesIO(put_call.kwargs["data"]))
    slide = uploaded.slides[0]
    assert slide.shapes.title.text == "New Title"
    assert slide.placeholders[1].text_frame.text == "New Body"
    assert result["title_applied"] is True
    assert result["body_applied"] is True


def test_add_slide_reports_unrequested_content_as_not_applied(monkeypatch):
    content = _pptx_bytes()
    _mock_versioned_write(monkeypatch, content)

    result = json.loads(powerpoint.powerpoint_add_slide("Deck.pptx", '"etag-1"'))

    assert result["status"] == "success"
    assert result["title_applied"] is False
    assert result["body_applied"] is False


def test_committed_add_slide_result_survives_low_platform_output_cap(monkeypatch):
    content = _pptx_bytes()
    _, _, mock_put = _mock_versioned_write(monkeypatch, content)
    mock_put.return_value = MockResponse(
        {
            "id": "item-1",
            "name": "Deck.pptx",
            "description": "provider metadata " * 100,
            "thumbnails": [{"large": {"url": "https://example.test/" + "x" * 500}}],
        },
        status_code=200,
    )
    max_chars = 80
    monkeypatch.setattr(powerpoint, "get_tool_max_output_length", lambda: max_chars)

    raw = powerpoint.powerpoint_add_slide("Deck.pptx", '"etag-1"')
    filtered = OutputValueFilter(
        max_chars=max_chars, max_fields=100, max_recursion=10
    ).filter(raw, "powerpoint_add_slide")

    assert filtered == raw
    assert len(raw) <= max_chars
    assert json.loads(filtered)["status"] == "success"


def test_add_slide_rejects_out_of_range_layout(monkeypatch):
    _mock_download(monkeypatch, _pptx_bytes())
    mock_session_cls = MagicMock()
    monkeypatch.setattr(powerpoint.requests, "Session", mock_session_cls)

    result = json.loads(
        powerpoint.powerpoint_add_slide("Deck.pptx", '"etag-1"', layout_index=999)
    )
    assert result["status"] == "error"
    assert result["message"].startswith("layout_index 999 is out of range")
    mock_session_cls.assert_not_called()


def test_add_slide_rejects_stale_etag_before_download(monkeypatch):
    content = _pptx_bytes()
    mock_request = Mock(return_value=_metadata_response(content, '"etag-new"'))
    mock_get = Mock()
    monkeypatch.setattr(powerpoint.requests, "request", mock_request)
    monkeypatch.setattr(powerpoint.requests, "get", mock_get)

    result = json.loads(powerpoint.powerpoint_add_slide("Deck.pptx", '"etag-old"'))

    assert result["status"] == "conflict"
    mock_get.assert_not_called()


@pytest.mark.parametrize(
    ("kwargs", "missing_field"),
    [
        ({"title": "Not applicable", "layout_index": 6}, "title"),
        ({"body_text": "Not applicable", "layout_index": 6}, "body_text"),
    ],
)
def test_add_slide_rejects_content_missing_from_layout(
    monkeypatch, kwargs, missing_field
):
    mock_request, _ = _mock_download(monkeypatch, _pptx_bytes())

    result = json.loads(
        powerpoint.powerpoint_add_slide("Deck.pptx", '"etag-1"', **kwargs)
    )

    assert result["status"] == "error"
    assert missing_field in result["message"]
    assert mock_request.call_count == 1


def test_add_slide_rejects_non_int_layout_index(monkeypatch):
    mock_request = Mock()
    monkeypatch.setattr(powerpoint.requests, "request", mock_request)

    result = json.loads(
        powerpoint.powerpoint_add_slide("Deck.pptx", '"etag-1"', layout_index="1")
    )

    assert result["status"] == "error"
    assert "layout_index must be an integer" in result["message"]
    mock_request.assert_not_called()


def test_set_shape_text_success(monkeypatch):
    def build(prs):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "Old Title"

    content = _pptx_bytes(build)
    _, _, mock_put = _mock_versioned_write(monkeypatch, content)

    result = json.loads(
        powerpoint.powerpoint_set_shape_text("Deck.pptx", 0, 0, "New Title", '"etag-1"')
    )

    assert result["status"] == "success"
    put_call = mock_put.call_args
    uploaded = Presentation(io.BytesIO(put_call.kwargs["data"]))
    assert uploaded.slides[0].shapes.title.text == "New Title"


def test_set_shape_text_rejects_non_int_indices(monkeypatch):
    mock_request = Mock()
    monkeypatch.setattr(powerpoint.requests, "request", mock_request)

    result = json.loads(
        powerpoint.powerpoint_set_shape_text("Deck.pptx", "0", 0, "text", '"etag-1"')
    )
    assert result["status"] == "error"
    assert "slide_index must be an integer" in result["message"]

    result = json.loads(
        powerpoint.powerpoint_set_shape_text("Deck.pptx", 0, "0", "text", '"etag-1"')
    )
    assert result["status"] == "error"
    assert "shape_index must be an integer" in result["message"]

    mock_request.assert_not_called()


def test_set_shape_text_rejects_shape_without_text_frame(monkeypatch):
    def build(prs):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        from pptx.enum.shapes import MSO_CONNECTOR

        slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, Inches(0), Inches(0), Inches(1), Inches(1)
        )

    content = _pptx_bytes(build)
    _mock_download(monkeypatch, content)

    result = json.loads(
        powerpoint.powerpoint_set_shape_text("Deck.pptx", 0, 0, "text", '"etag-1"')
    )

    assert result["status"] == "error"
    assert "no text frame" in result["message"]


def test_text_frame_has_dynamic_content_detects_hyperlink():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    run = box.text_frame.paragraphs[0].add_run()
    run.text = "Click me"
    run.hyperlink.address = "https://example.com"

    assert powerpoint._text_frame_has_dynamic_content(box.text_frame)


def test_text_frame_has_dynamic_content_detects_mouseover_hyperlink():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    run = box.text_frame.paragraphs[0].add_run()
    run.text = "Hover me"
    run_properties = run._r.get_or_add_rPr()
    etree.SubElement(run_properties, qn("a:hlinkMouseOver"))

    assert powerpoint._text_frame_has_dynamic_content(box.text_frame)


def test_text_frame_distinct_run_formatting_is_detected():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    bold_run = box.text_frame.paragraphs[0].add_run()
    bold_run.text = "Bold"
    bold_run.font.bold = True
    box.text_frame.paragraphs[0].add_run().text = "Plain"

    assert powerpoint._text_frame_has_distinct_run_formatting(box.text_frame)


def test_text_frame_equivalent_run_formatting_is_not_rejected():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    for text in ("First", "Second"):
        run = box.text_frame.paragraphs[0].add_run()
        run.text = text
        run.font.bold = True
        run.font.size = Pt(24)

    assert not powerpoint._text_frame_has_distinct_run_formatting(box.text_frame)


def test_text_frame_distinct_line_break_formatting_is_detected():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    box.text_frame.text = "First\vSecond"
    for run in box.text_frame.paragraphs[0].runs:
        run.font.size = Pt(24)
    line_break = box.text_frame.paragraphs[0]._p.find(qn("a:br"))
    line_break.get_or_add_rPr().set("sz", "4000")

    assert powerpoint._text_frame_has_distinct_run_formatting(box.text_frame)


def test_text_frame_has_dynamic_content_false_for_plain_text():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    box.text_frame.text = "Plain text"

    assert not powerpoint._text_frame_has_dynamic_content(box.text_frame)


def test_set_shape_text_rejects_shape_with_hyperlink(monkeypatch):
    def build(prs):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
        run = box.text_frame.paragraphs[0].add_run()
        run.text = "Click me"
        run.hyperlink.address = "https://example.com"

    content = _pptx_bytes(build)
    _mock_download(monkeypatch, content)

    result = json.loads(
        powerpoint.powerpoint_set_shape_text("Deck.pptx", 0, 0, "New text", '"etag-1"')
    )

    assert result["status"] == "error"
    assert "hyperlink" in result["message"]


def test_delete_slide_tool_uploads_updated_presentation(monkeypatch):
    def build(prs):
        prs.slides.add_slide(prs.slide_layouts[1]).shapes.title.text = "A"
        prs.slides.add_slide(prs.slide_layouts[1]).shapes.title.text = "B"

    content = _pptx_bytes(build)
    _, _, mock_put = _mock_versioned_write(monkeypatch, content)

    result = json.loads(powerpoint.powerpoint_delete_slide("Deck.pptx", 0, '"etag-1"'))

    assert result["status"] == "success"
    put_call = mock_put.call_args
    uploaded = Presentation(io.BytesIO(put_call.kwargs["data"]))
    assert [s.shapes.title.text for s in uploaded.slides] == ["B"]


def test_delete_slide_tool_rejects_non_int_slide_index(monkeypatch):
    mock_request = Mock()
    monkeypatch.setattr(powerpoint.requests, "request", mock_request)

    result = json.loads(
        powerpoint.powerpoint_delete_slide("Deck.pptx", "0", '"etag-1"')
    )

    assert result["status"] == "error"
    assert "slide_index must be an integer" in result["message"]
    mock_request.assert_not_called()


def test_upload_presentation_rejects_oversized_content(monkeypatch):
    presentation = Presentation()
    monkeypatch.setattr(powerpoint, "_MAX_PRESENTATION_BYTES", 10)

    with pytest.raises(ValueError, match="MB limit"):
        powerpoint._upload_presentation(
            presentation, "Deck.pptx", None, None, '"etag-1"'
        )


def test_upload_presentation_conditionally_creates_upload_session(monkeypatch):
    """Every replacement uses If-Match when its upload session is created."""
    presentation = Presentation()
    monkeypatch.setattr(powerpoint, "_UPLOAD_SESSION_CHUNK_BYTES", 10)
    mock_request = Mock(
        return_value=MockResponse({"uploadUrl": "https://upload.example/session"})
    )
    monkeypatch.setattr(powerpoint.requests, "request", mock_request)
    put_calls: list[MockResponse] = []

    def fake_put(url, *, data, headers, timeout):
        put_calls.append((url, data, headers))
        range_part, total_str = (
            headers["Content-Range"].removeprefix("bytes ").split("/")
        )
        _, end_str = range_part.split("-")
        is_last = int(end_str) + 1 == int(total_str)
        if is_last:
            return MockResponse({"id": "item-1"}, status_code=200)
        return MockResponse(
            {"nextExpectedRanges": [f"{int(end_str) + 1}-"]}, status_code=202
        )

    mock_session_put = Mock(side_effect=fake_put)
    mock_session_cls = MagicMock()
    mock_session_cls.return_value.__enter__.return_value.put = mock_session_put
    monkeypatch.setattr(powerpoint.requests, "Session", mock_session_cls)

    item = powerpoint._upload_presentation(
        presentation, "Deck.pptx", None, None, '"etag-1"'
    )

    assert item["id"] == "item-1"
    assert len(put_calls) > 1
    assert mock_request.call_args.kwargs["headers"]["If-Match"] == '"etag-1"'
    assert "Authorization" not in put_calls[0][2]
    full_content = b"".join(call[1] for call in put_calls)
    Presentation(io.BytesIO(full_content))


def test_upload_presentation_session_rejects_when_no_upload_url(monkeypatch):
    mock_request = Mock(return_value=MockResponse({}))
    monkeypatch.setattr(powerpoint.requests, "request", mock_request)

    with pytest.raises(RuntimeError, match="did not return an upload session URL"):
        powerpoint._upload_presentation_session(
            b"content", "Deck.pptx", None, None, '"etag-1"'
        )


def test_upload_presentation_session_failure_does_not_leak_session_url(monkeypatch):
    secret_url = "https://upload.example/session?token=super-secret-value"
    mock_request = Mock(return_value=MockResponse({"uploadUrl": secret_url}))
    monkeypatch.setattr(powerpoint.requests, "request", mock_request)
    mock_put = Mock(return_value=MockResponse({}, status_code=500, url=secret_url))
    mock_session_cls = MagicMock()
    mock_http = mock_session_cls.return_value.__enter__.return_value
    mock_http.put = mock_put
    monkeypatch.setattr(powerpoint.requests, "Session", mock_session_cls)

    with pytest.raises(powerpoint._IndeterminateWriteError) as exc_info:
        powerpoint._upload_presentation_session(
            b"content" * 100_000, "Deck.pptx", None, None, '"etag-1"'
        )

    assert exc_info.value.__cause__ is None
    assert "super-secret-value" not in str(exc_info.value)
    mock_http.delete.assert_not_called()


def test_upload_session_cancels_after_definite_nonfinal_failure(monkeypatch):
    secret_url = "https://upload.example/session?token=super-secret-value"
    monkeypatch.setattr(powerpoint, "_UPLOAD_SESSION_CHUNK_BYTES", 5)
    monkeypatch.setattr(
        powerpoint.requests,
        "request",
        Mock(return_value=MockResponse({"uploadUrl": secret_url})),
    )
    responses = iter(
        [
            MockResponse({"nextExpectedRanges": ["5-"]}, status_code=202),
            MockResponse({}, status_code=400, url=secret_url),
        ]
    )
    mock_session_cls = MagicMock()
    mock_http = mock_session_cls.return_value.__enter__.return_value
    mock_http.put = Mock(side_effect=lambda *a, **k: next(responses))
    monkeypatch.setattr(powerpoint.requests, "Session", mock_session_cls)

    with pytest.raises(powerpoint._GraphRequestError) as exc_info:
        powerpoint._upload_presentation_session(
            b"0123456789", "Deck.pptx", None, None, '"etag-1"'
        )

    assert "super-secret-value" not in str(exc_info.value)
    mock_http.delete.assert_called_once_with(
        secret_url, timeout=powerpoint._UPLOAD_CANCEL_TIMEOUT_SECONDS
    )


def test_upload_session_cancels_after_invalid_progress(monkeypatch):
    upload_url = "https://upload.example/session"
    monkeypatch.setattr(powerpoint, "_UPLOAD_SESSION_CHUNK_BYTES", 5)
    monkeypatch.setattr(
        powerpoint.requests,
        "request",
        Mock(return_value=MockResponse({"uploadUrl": upload_url})),
    )
    mock_session_cls = MagicMock()
    mock_http = mock_session_cls.return_value.__enter__.return_value
    mock_http.put = Mock(
        return_value=MockResponse({"nextExpectedRanges": []}, status_code=202)
    )
    monkeypatch.setattr(powerpoint.requests, "Session", mock_session_cls)

    with pytest.raises(RuntimeError, match="empty.*progress"):
        powerpoint._upload_presentation_session(
            b"0123456789", "Deck.pptx", None, None, '"etag-1"'
        )

    mock_http.delete.assert_called_once_with(
        upload_url, timeout=powerpoint._UPLOAD_CANCEL_TIMEOUT_SECONDS
    )


def test_upload_session_follows_server_reported_next_offset(monkeypatch):
    monkeypatch.setattr(powerpoint, "_UPLOAD_SESSION_CHUNK_BYTES", 5)
    mock_request = Mock(
        return_value=MockResponse({"uploadUrl": "https://upload.example/session"})
    )
    monkeypatch.setattr(powerpoint.requests, "request", mock_request)
    responses = iter(
        [
            MockResponse({"nextExpectedRanges": ["3-"]}, status_code=202),
            MockResponse({"nextExpectedRanges": ["8-"]}, status_code=202),
            MockResponse({"id": "item-1"}, status_code=200),
        ]
    )
    mock_put = Mock(side_effect=lambda *a, **k: next(responses))
    mock_session_cls = MagicMock()
    mock_session_cls.return_value.__enter__.return_value.put = mock_put
    monkeypatch.setattr(powerpoint.requests, "Session", mock_session_cls)

    item = powerpoint._upload_presentation_session(
        b"0123456789", "Deck.pptx", None, None, '"etag-1"'
    )

    assert item["id"] == "item-1"
    assert [
        call.kwargs["headers"]["Content-Range"] for call in mock_put.call_args_list
    ] == ["bytes 0-4/10", "bytes 3-7/10", "bytes 8-9/10"]


def test_add_slide_maps_upload_precondition_failure_to_conflict(monkeypatch):
    content = _pptx_bytes()
    responses = iter(
        [
            _metadata_response(content),
            MockResponse({"error": {"code": "preconditionFailed"}}, status_code=412),
        ]
    )
    monkeypatch.setattr(
        powerpoint.requests,
        "request",
        Mock(side_effect=lambda *a, **k: next(responses)),
    )
    monkeypatch.setattr(
        powerpoint.requests, "get", Mock(return_value=MockResponse(content=content))
    )

    result = json.loads(powerpoint.powerpoint_add_slide("Deck.pptx", '"etag-1"'))

    assert result["status"] == "conflict"


def test_add_slide_maps_final_fragment_precondition_failure_to_conflict(monkeypatch):
    content = _pptx_bytes()
    _, _, mock_put = _mock_versioned_write(monkeypatch, content)
    mock_put.return_value = MockResponse(
        {"error": {"code": "preconditionFailed"}}, status_code=412
    )

    result = json.loads(powerpoint.powerpoint_add_slide("Deck.pptx", '"etag-1"'))

    assert result["status"] == "conflict"
    assert "changed" in result["message"]


def test_add_slide_reports_final_connection_loss_as_indeterminate(monkeypatch):
    content = _pptx_bytes()
    _, _, mock_put = _mock_versioned_write(monkeypatch, content)
    mock_put.side_effect = requests.ConnectionError(
        "lost response from https://upload.example/session?secret=token"
    )

    result = json.loads(powerpoint.powerpoint_add_slide("Deck.pptx", '"etag-1"'))

    assert result["status"] == "indeterminate"
    assert result["safe_to_retry"] is False
    assert "secret=token" not in result["message"]


# ---------------------------------------------------------------------------
# error propagation
# ---------------------------------------------------------------------------


def test_graph_request_http_error_does_not_leak_redirect_url(monkeypatch):
    """Graph's /content GET redirects to a preauthenticated download URL --
    if that final request fails, requests' own HTTPError.__str__ embeds
    that signed URL as response.url. _graph_request must never surface
    that string, whether directly in the raised message or as the
    chained exception's __cause__ (a future traceback/log/APM capture
    could still surface a __cause__ even with a clean top-level message)."""
    secret_url = "https://blob.example/deck.pptx?token=super-secret-value"
    mock_request = Mock(return_value=MockResponse({}, status_code=403, url=secret_url))
    monkeypatch.setattr(powerpoint.requests, "request", mock_request)

    with pytest.raises(powerpoint._GraphRequestError) as exc_info:
        powerpoint._graph_request("GET", "/me/drive/root:/Deck.pptx:/content")

    assert "super-secret-value" not in str(exc_info.value)
    assert secret_url not in str(exc_info.value)
    assert exc_info.value.__cause__ is None


def test_graph_request_does_not_forward_secret_from_error_body(monkeypatch):
    secret_url = "https://blob.example/deck.pptx?sig=super-secret-value"
    response = MockResponse(
        {"error": {"code": "accessDenied", "message": secret_url}},
        status_code=403,
    )
    monkeypatch.setattr(powerpoint.requests, "request", Mock(return_value=response))

    with pytest.raises(powerpoint._GraphRequestError) as exc_info:
        powerpoint._graph_request("GET", "/me/drive/root:/Deck.pptx:/content")

    assert "accessDenied" in str(exc_info.value)
    assert "super-secret-value" not in str(exc_info.value)


def test_graph_request_connection_error_does_not_leak_redirect_url(monkeypatch):
    """The initial requests.request() call itself (not just raise_for_status())
    can fail at the connection layer (timeout, reset, TLS error) after
    following the same redirect -- requests.RequestException's own str()
    embeds the URL just like HTTPError's does, so this path must be
    sanitized too."""
    secret_url = "https://blob.example/deck.pptx?token=super-secret-value"
    mock_request = Mock(
        side_effect=requests.ConnectionError(f"Connection refused: {secret_url}")
    )
    monkeypatch.setattr(powerpoint.requests, "request", mock_request)

    with pytest.raises(RuntimeError) as exc_info:
        powerpoint._graph_request("GET", "/me/drive/root:/Deck.pptx:/content")

    assert "super-secret-value" not in str(exc_info.value)
    assert exc_info.value.__cause__ is None


def test_get_presentation_text_reports_download_failure_without_leaking_url(
    monkeypatch,
):
    secret_url = "https://blob.example/deck.pptx?token=super-secret-value"
    mock_request = Mock(return_value=MockResponse({}, status_code=403, url=secret_url))
    monkeypatch.setattr(powerpoint.requests, "request", mock_request)

    result = json.loads(powerpoint.powerpoint_get_presentation_text("Deck.pptx"))

    assert result["status"] == "error"
    assert "super-secret-value" not in result["message"]
    assert secret_url not in result["message"]


def test_slide_notes_returns_none_when_notes_placeholder_is_missing():
    """notes_text_frame can be None even when has_notes_slide is True --
    e.g. the notes placeholder was deleted from the notes slide while the
    notes-slide XML part itself survives."""
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    notes_slide = slide.notes_slide  # creates the notes slide
    notes_placeholder = notes_slide.notes_placeholder
    notes_placeholder._element.getparent().remove(notes_placeholder._element)

    assert notes_slide.notes_text_frame is None
    assert powerpoint._slide_notes(slide) is None


def _non_json_success_response() -> Mock:
    """A 2xx response whose body isn't valid JSON -- unlike MockResponse,
    whose .json() always returns a stored dict regardless of `content`,
    this actually raises like a real empty/non-JSON body would."""
    response = Mock()
    response.status_code = 200
    response.raise_for_status = Mock(return_value=None)
    response.json = Mock(side_effect=ValueError("Expecting value"))
    return response


def test_upload_presentation_session_marks_non_json_final_response_indeterminate(
    monkeypatch,
):
    """A 2xx final response without an item may follow a committed write."""
    mock_request = Mock(
        return_value=MockResponse({"uploadUrl": "https://upload.example/session"})
    )
    monkeypatch.setattr(powerpoint.requests, "request", mock_request)
    mock_put = Mock(return_value=_non_json_success_response())
    mock_session_cls = MagicMock()
    mock_session_cls.return_value.__enter__.return_value.put = mock_put
    monkeypatch.setattr(powerpoint.requests, "Session", mock_session_cls)

    with pytest.raises(powerpoint._IndeterminateWriteError, match="did not return"):
        powerpoint._upload_presentation_session(
            b"content", "Deck.pptx", None, None, '"etag-1"'
        )


def test_create_only_upload_handles_non_json_final_response(monkeypatch):
    mock_request = Mock(
        return_value=MockResponse({"uploadUrl": "https://upload.example/session"})
    )
    monkeypatch.setattr(powerpoint.requests, "request", mock_request)
    mock_put = Mock(return_value=_non_json_success_response())
    monkeypatch.setattr(powerpoint.requests, "put", mock_put)

    with pytest.raises(powerpoint._IndeterminateWriteError, match="did not confirm"):
        powerpoint._create_only_upload(b"content", "Deck.pptx", None, None)


def test_create_presentation_reports_non_json_final_response_as_indeterminate(
    monkeypatch,
):
    mock_request = Mock(
        return_value=MockResponse({"uploadUrl": "https://upload.example/session"})
    )
    monkeypatch.setattr(powerpoint.requests, "request", mock_request)
    monkeypatch.setattr(
        powerpoint.requests, "put", Mock(return_value=_non_json_success_response())
    )

    result = json.loads(powerpoint.powerpoint_create_presentation("Deck.pptx"))

    assert result["status"] == "indeterminate"
    assert result["safe_to_retry"] is False
    assert "check whether the file exists" in result["message"]


def test_replace_text_frame_text_preserves_run_formatting():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    run = box.text_frame.paragraphs[0].add_run()
    run.text = "Old"
    run.font.bold = True
    run.font.size = Pt(24)

    powerpoint._replace_text_frame_text(box.text_frame, "New")

    new_run = box.text_frame.paragraphs[0].runs[0]
    assert new_run.text == "New"
    assert new_run.font.bold is True
    assert new_run.font.size == Pt(24)


def test_replace_text_frame_text_preserves_paragraph_alignment():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    box.text_frame.text = "Old"
    box.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

    powerpoint._replace_text_frame_text(box.text_frame, "New")

    assert box.text_frame.paragraphs[0].alignment == PP_ALIGN.CENTER
    assert box.text_frame.paragraphs[0].text == "New"


def test_replace_text_frame_text_handles_paragraph_count_change():
    """New paragraphs inherit the final old paragraph's formatting."""
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    box.text_frame.text = "Only line"
    box.text_frame.paragraphs[0].runs[0].font.bold = True
    box.text_frame.paragraphs[0].runs[0].font.size = Pt(24)
    box.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

    powerpoint._replace_text_frame_text(box.text_frame, "Line one\nLine two")

    paragraphs = box.text_frame.paragraphs
    assert [p.text for p in paragraphs] == ["Line one", "Line two"]
    for paragraph in paragraphs:
        assert paragraph.alignment == PP_ALIGN.CENTER
        assert paragraph.runs[0].font.bold is True
        assert paragraph.runs[0].font.size == Pt(24)


def test_replace_text_frame_text_preserves_formatting_across_soft_break():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    run = box.text_frame.paragraphs[0].add_run()
    run.text = "Old"
    run.font.bold = True
    run.font.size = Pt(24)

    powerpoint._replace_text_frame_text(box.text_frame, "Line one\vLine two")

    runs = box.text_frame.paragraphs[0].runs
    assert [run.text for run in runs] == ["Line one", "Line two"]
    for result_run in runs:
        assert result_run.font.bold is True
        assert result_run.font.size == Pt(24)
    line_break = box.text_frame.paragraphs[0]._p.find(qn("a:br"))
    line_break_properties = line_break.find(qn("a:rPr"))
    assert line_break_properties is not None
    assert line_break_properties.get("b") == "1"
    assert line_break_properties.get("sz") == "2400"


def test_replace_text_frame_text_preserves_end_paragraph_formatting():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    box.text_frame.text = "Old"
    end_properties = box.text_frame.paragraphs[0]._p.get_or_add_endParaRPr()
    end_properties.set("b", "1")
    end_properties.set("sz", "3200")

    powerpoint._replace_text_frame_text(box.text_frame, "New")

    result = box.text_frame.paragraphs[0]._p.find(qn("a:endParaRPr"))
    assert result is not None
    assert result.get("b") == "1"
    assert result.get("sz") == "3200"


def test_set_shape_text_preserves_formatting(monkeypatch):
    def build(prs):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
        run = box.text_frame.paragraphs[0].add_run()
        run.text = "Old"
        run.font.bold = True

    content = _pptx_bytes(build)
    _, _, mock_put = _mock_versioned_write(monkeypatch, content)

    result = json.loads(
        powerpoint.powerpoint_set_shape_text("Deck.pptx", 0, 0, "New", '"etag-1"')
    )

    assert result["status"] == "success"
    put_call = mock_put.call_args
    uploaded = Presentation(io.BytesIO(put_call.kwargs["data"]))
    run = uploaded.slides[0].shapes[0].text_frame.paragraphs[0].runs[0]
    assert run.text == "New"
    assert run.font.bold is True


def test_set_shape_text_preserves_formatting_across_structural_breaks(monkeypatch):
    def build(prs):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
        paragraph = box.text_frame.paragraphs[0]
        paragraph.alignment = PP_ALIGN.CENTER
        run = paragraph.add_run()
        run.text = "Old"
        run.font.bold = True
        run.font.size = Pt(24)
        end_properties = paragraph._p.get_or_add_endParaRPr()
        end_properties.set("lang", "en-US")

    content = _pptx_bytes(build)
    _, _, mock_put = _mock_versioned_write(monkeypatch, content)

    result = json.loads(
        powerpoint.powerpoint_set_shape_text(
            "Deck.pptx", 0, 0, "Line one\nLine two\vLine three", '"etag-1"'
        )
    )

    assert result["status"] == "success"
    uploaded = Presentation(io.BytesIO(mock_put.call_args.kwargs["data"]))
    paragraphs = uploaded.slides[0].shapes[0].text_frame.paragraphs
    assert [paragraph.text for paragraph in paragraphs] == [
        "Line one",
        "Line two\vLine three",
    ]
    for paragraph in paragraphs:
        assert paragraph.alignment == PP_ALIGN.CENTER
        for run in paragraph.runs:
            assert run.font.bold is True
            assert run.font.size == Pt(24)
        end_properties = paragraph._p.find(qn("a:endParaRPr"))
        assert end_properties is not None
        assert end_properties.get("lang") == "en-US"
    line_break = paragraphs[1]._p.find(qn("a:br"))
    line_break_properties = line_break.find(qn("a:rPr"))
    assert line_break_properties is not None
    assert line_break_properties.get("b") == "1"
    assert line_break_properties.get("sz") == "2400"


def test_set_shape_text_allows_equivalently_formatted_runs(monkeypatch):
    def build(prs):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
        for value in ("Old ", "text"):
            run = box.text_frame.paragraphs[0].add_run()
            run.text = value
            run.font.bold = True
            run.font.size = Pt(24)

    content = _pptx_bytes(build)
    _, _, mock_put = _mock_versioned_write(monkeypatch, content)

    result = json.loads(
        powerpoint.powerpoint_set_shape_text("Deck.pptx", 0, 0, "New", '"etag-1"')
    )

    assert result["status"] == "success"
    uploaded = Presentation(io.BytesIO(mock_put.call_args.kwargs["data"]))
    runs = uploaded.slides[0].shapes[0].text_frame.paragraphs[0].runs
    assert len(runs) == 1
    assert runs[0].text == "New"
    assert runs[0].font.bold is True
    assert runs[0].font.size == Pt(24)


def test_missing_auth_token_is_reported(monkeypatch):
    monkeypatch.delenv("AUTH_TOKEN", raising=False)

    result = json.loads(powerpoint.powerpoint_get_presentation_text("Deck.pptx"))

    assert result["status"] == "error"
    assert "AUTH_TOKEN" in result["message"]
