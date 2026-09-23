import base64
import binascii
import io
import json
import logging
import os
import zipfile
from copy import deepcopy
from dataclasses import dataclass
from typing import Any, Callable
from urllib.parse import quote

import requests
from mcp.server.fastmcp import FastMCP
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.presentation import Presentation as PresentationType

from ....config import get_tool_max_output_length
from ....core.tools.core.file_analysis import iter_pptx_shapes
from .utils import setup_proxy_env, url_path_id

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger("powerpoint-mcp")

setup_proxy_env()

mcp = FastMCP("powerpoint-mcp")

GRAPH_BASE_URL = "https://graph.microsoft.com/v1.0"
DEFAULT_TIMEOUT_SECONDS = 30
# Matches onedrive.py's _BINARY_UPLOAD_TIMEOUT_SECONDS for the same class of
# operation (a large binary GET/PUT); used for both directions here since
# this module's download and upload can each move up to
# _MAX_PRESENTATION_BYTES.
_BINARY_TIMEOUT_SECONDS = 120

_POWERPOINT_MIME_TYPE = (
    "application/vnd.openxmlformats-officedocument.presentationml.presentation"
)

# Microsoft Graph exposes a PowerPoint (.pptx) file only as a driveItem
# content blob -- there is no structured "PowerPoint API" resource (no
# per-slide/per-shape Graph endpoints). Every tool here downloads the whole
# file, edits it in memory with python-pptx, and re-uploads the whole file.
# Kept entirely in memory (io.BytesIO, never written to local disk).
#
# A deliberately arbitrary product ceiling (independent of any Graph-side
# limit) on a presentation's compressed size -- unlike
# onedrive.py's file upload, which streams a chunk at a time straight from
# disk, this module's whole download-edit-reupload cycle holds the
# serialized bytes plus python-pptx's object graph in memory.
_MAX_PRESENTATION_BYTES = 200_000_000

# A small .pptx can still be a zip bomb. Bound the total declared size of
# OOXML members before python-pptx expands them into objects.
_MAX_PRESENTATION_UNCOMPRESSED_BYTES = 500_000_000
_MAX_PRESENTATION_PARTS = 10_000
_DOWNLOAD_CHUNK_BYTES = 1024 * 1024

# Graph requires every upload-session fragment but the last to be a
# multiple of 320 KiB, and documents a 60 MiB hard maximum per PUT; 5 MiB
# is also in its recommended 5-10 MiB "best practice" range, and matches
# onedrive.py's own chunk size for the same API.
_UPLOAD_SESSION_CHUNK_BYTES = 5 * 1024 * 1024
_UPLOAD_CANCEL_TIMEOUT_SECONDS = 10

# python-pptx has no public API for adding a slide at other than the layout
# picked, or for removing one at all -- see _delete_slide's own docstring
# for the latter. Layout 1 ("Title and Content") is the conventional
# default for "just add a slide" the same way Word's own UI defaults a new
# document to the Normal style.
_DEFAULT_SLIDE_LAYOUT_INDEX = 1


class _GraphRequestError(RuntimeError):
    """Graph HTTP failure that retains its status without response parsing."""

    def __init__(self, message: str, *, status_code: int) -> None:
        super().__init__(message)
        self.status_code = status_code


class _ConflictError(RuntimeError):
    """The caller's presentation version is no longer current."""


class _IndeterminateWriteError(RuntimeError):
    """Graph may have committed a write whose final response was lost."""


@dataclass(frozen=True)
class _PresentationSnapshot:
    presentation: PresentationType
    etag: str


def _compact_json(payload: dict[str, Any]) -> str:
    return json.dumps(payload, ensure_ascii=False, separators=(",", ":"))


def _bounded_envelope(
    candidates: list[dict[str, Any]], *, fallback: dict[str, Any]
) -> str:
    """Choose the richest valid JSON envelope that the string filter won't cut."""
    max_chars = get_tool_max_output_length()
    for payload in candidates:
        response = _compact_json(payload)
        if len(response) <= max_chars:
            return response
    # An operator can configure a cap below even the smallest contract-bearing
    # object. Preserve valid JSON and the outcome status rather than pretending
    # an empty object is a usable mutation result.
    return _compact_json(fallback)


def _caller_safe_item(item: Any) -> dict[str, Any]:
    """Keep only stable, scalar driveItem fields useful after a mutation."""
    if not isinstance(item, dict):
        return {}
    return {
        field: value
        for field in ("id", "name", "eTag", "cTag", "size", "webUrl")
        if isinstance((value := item.get(field)), (str, int, float, bool))
    }


def _success(**payload: Any) -> str:
    compact_payload = dict(payload)
    if "item" in compact_payload:
        compact_payload["item"] = _caller_safe_item(compact_payload["item"])
    without_item = {
        key: value for key, value in compact_payload.items() if key != "item"
    }
    return _bounded_envelope(
        [
            {"status": "success", **payload},
            {"status": "success", **compact_payload},
            {"status": "success", **without_item},
            {"status": "success"},
        ],
        fallback={"status": "success"},
    )


def _error(message: str, *, details: Any = None) -> str:
    candidates: list[dict[str, Any]] = []
    if details is not None:
        candidates.append({"status": "error", "message": message, "details": details})
    candidates.extend(
        [
            {"status": "error", "message": message},
            {"status": "error", "message": "output cap too small"},
            {"status": "error"},
        ]
    )
    return _bounded_envelope(candidates, fallback={"status": "error"})


def _conflict(message: str) -> str:
    return _bounded_envelope(
        [{"status": "conflict", "message": message}, {"status": "conflict"}],
        fallback={"status": "conflict"},
    )


def _indeterminate(message: str) -> str:
    minimal = {"status": "indeterminate", "safe_to_retry": False}
    return _bounded_envelope(
        [{**minimal, "message": message}, minimal],
        fallback=minimal,
    )


def _bounded_error(message: str, *, details: Any = None) -> str:
    """Backward-compatible name used by bounded read response builders."""
    return _error(message, details=details)


def _encode_read_cursor(index: int, etag: str, scope: str) -> str:
    payload = _compact_json({"index": index, "etag": etag, "scope": scope})
    return base64.urlsafe_b64encode(payload.encode()).decode().rstrip("=")


def _decode_read_cursor(
    cursor: str | None, *, etag: str, scope: str, total_count: int
) -> int:
    if cursor is None:
        return 0
    if not isinstance(cursor, str) or not cursor:
        raise ValueError("cursor must be a non-empty string")
    try:
        padding = "=" * (-len(cursor) % 4)
        payload = json.loads(base64.urlsafe_b64decode(cursor + padding).decode())
        index = payload["index"]
        cursor_etag = payload["etag"]
        cursor_scope = payload["scope"]
    except (
        binascii.Error,
        UnicodeDecodeError,
        json.JSONDecodeError,
        KeyError,
        TypeError,
    ) as exc:
        raise ValueError("cursor is invalid") from exc
    if cursor_etag != etag:
        raise _ConflictError(
            "The presentation changed while reading a paginated result; restart "
            "without a cursor"
        )
    if cursor_scope != scope:
        raise ValueError("cursor does not belong to this PowerPoint read operation")
    if not isinstance(index, int) or isinstance(index, bool):
        raise ValueError("cursor is invalid")
    if not 0 <= index <= total_count:
        raise ValueError("cursor is out of range")
    return index


def _bounded_page_response(
    *,
    field_name: str,
    total_count: int,
    start_index: int,
    item_at: Callable[[int], Any],
    item_label: str,
    etag: str,
    scope: str,
) -> str:
    """Serialize one resumable page without relying on destructive filtering.

    MCP results are transported as a single string and the platform applies a
    per-string output cap after the tool returns. Building the page here keeps
    that outer filter from cutting a JSON document in half. If one atomic item
    cannot fit, return an explicit ``omitted_item`` record and a cursor past it
    instead of stranding all later items behind an unrecoverable value.
    """
    max_chars = get_tool_max_output_length()

    def render(items: list[Any], next_index: int) -> str:
        truncated = next_index < total_count
        return _compact_json(
            {
                "status": "success",
                field_name: items,
                "etag": etag,
                "truncated": truncated,
                "next_cursor": (
                    _encode_read_cursor(next_index, etag, scope) if truncated else None
                ),
                "total_count": total_count,
            }
        )

    items: list[Any] = []
    next_index = start_index
    empty_page = render(items, next_index)
    if len(empty_page) > max_chars:
        return _bounded_error(
            "PowerPoint pagination metadata exceeds the configured output limit"
        )

    while next_index < total_count:
        item = item_at(next_index)
        candidate = render([*items, item], next_index + 1)
        if len(candidate) > max_chars:
            if not items:
                # Do not strand every later item behind one value that can never
                # fit. Explicitly report the omission and issue a forward cursor;
                # the cursor is emitted even when this is the final item so the
                # caller can observe a terminal, non-truncated page next.
                omitted_index = next_index
                next_index += 1
                omitted_response = _compact_json(
                    {
                        "status": "success",
                        field_name: [],
                        "etag": etag,
                        "truncated": True,
                        "next_cursor": _encode_read_cursor(next_index, etag, scope),
                        "total_count": total_count,
                        "omitted_item": {
                            "item_index": omitted_index,
                            "item_type": item_label,
                            "reason": "item_exceeds_output_limit",
                            "output_limit": max_chars,
                        },
                    }
                )
                if len(omitted_response) <= max_chars:
                    return omitted_response
                return _bounded_error(
                    "PowerPoint oversized-item metadata exceeds the configured "
                    "output limit"
                )
            break
        items.append(item)
        next_index += 1

    return render(items, next_index)


def _graph_headers(extra_headers: dict[str, str] | None = None) -> dict[str, str]:
    token = os.environ.get("AUTH_TOKEN")
    if not token:
        raise ValueError("AUTH_TOKEN environment variable is missing")

    headers = {
        "Authorization": f"Bearer {token}",
        "Accept": "application/json",
        "Content-Type": "application/json",
    }
    if extra_headers:
        headers.update(extra_headers)
    return headers


def _graph_request(
    method: str,
    path: str,
    *,
    params: dict[str, Any] | None = None,
    body: dict[str, Any] | None = None,
    extra_headers: dict[str, str] | None = None,
    timeout: int = DEFAULT_TIMEOUT_SECONDS,
) -> Any:
    # Never stringify request exceptions or forward raw response bodies.
    # Graph and its backing storage may place preauthenticated URLs in either;
    # those URLs are short-lived bearer secrets. Build errors only from this
    # function's caller-controlled method/path and a parsed Graph error code.
    try:
        response = requests.request(
            method=method,
            url=f"{GRAPH_BASE_URL}{path}",
            headers=_graph_headers(extra_headers),
            params=params,
            json=body,
            timeout=timeout,
        )
    except requests.RequestException:
        raise RuntimeError(f"Graph request failed: {method} {path}") from None

    try:
        response.raise_for_status()
    except requests.HTTPError:
        message = f"Graph {method} {path} failed with HTTP {response.status_code}"
        # Never forward the raw response body. /content redirects to a
        # preauthenticated URL, and storage error bodies are allowed to echo
        # that bearer URL. A Graph error code is useful and safe enough.
        try:
            payload = response.json()
            error = payload.get("error") if isinstance(payload, dict) else None
            code = error.get("code") if isinstance(error, dict) else None
        except (ValueError, TypeError):
            code = None
        if isinstance(code, str) and code:
            message = f"{message} ({code})"
        raise _GraphRequestError(message, status_code=response.status_code) from None

    if response.status_code == 204 or not response.content:
        return {}
    return response.json()


def _site_segment(site_id: str) -> str:
    """Percent-encode a caller-supplied Graph site identifier for
    interpolation into a URL path segment.

    A Graph site id is one of: the literal "root", a composite id
    ("hostname,spSiteId,spWebId"), or a "hostname:/server-relative-path"
    form. ':' and '/' stay unescaped because they're structural to the
    third shape, while a '.'/'..' segment is rejected outright -- standard
    HTTP client URL normalization could otherwise walk the request off
    "/sites/{id}/..." and onto a different Graph endpoint under the same
    OAuth token.
    """
    if not isinstance(site_id, str) or not site_id.strip():
        raise ValueError("site_id is required")
    value = site_id.strip()
    if any(segment in (".", "..", "") for segment in value.split("/")):
        raise ValueError(
            f"site_id must not contain '.', '..', or empty (e.g. '//') segments: "
            f"{site_id!r}"
        )
    return quote(value, safe=":/,")


def _normalize_relative_path(path: str) -> str:
    """Normalize a drive-relative file path for a root:/{path}: request URL,
    rejecting '.'/'..' segments, a trailing folder separator, and a filename
    ending in a period (Graph/SharePoint's backing storage can silently
    normalize a trailing dot away, so "Deck.pptx." could silently resolve
    to a real, different "Deck.pptx")."""
    value = path.strip().strip("/")
    if not value:
        raise ValueError("file_path is required")
    if path.strip().endswith("/"):
        raise ValueError(
            "file_path must include a filename, not end with a folder separator"
        )
    if "\\" in value:
        raise ValueError("file_path must use '/' separators and must not contain '\\'")
    if any(segment in (".", "..", "") for segment in value.split("/")):
        raise ValueError(
            f"file_path must not contain '.', '..', or empty (e.g. '//') segments: "
            f"{path!r}"
        )
    if value.rsplit("/", 1)[-1].endswith("."):
        raise ValueError(f"file_path filename must not end with a period: {path!r}")
    if not value.rsplit("/", 1)[-1].lower().endswith(".pptx"):
        raise ValueError(
            "file_path must name a .pptx presentation; macro-enabled and other "
            "PowerPoint package types are not supported"
        )
    return value


def _item_path(file_path: str, site_id: str | None, drive_id: str | None) -> str:
    # Checked via "is not None" rather than truthiness: an empty string is
    # a caller mistake (e.g. an upstream field that defaults unset to ""
    # rather than None), not "not provided" -- a bare truthiness check
    # would silently treat it the same as None and fall through to
    # /me/drive, reading/editing the wrong drive with no indication site_id
    # or drive_id was ever ignored. This way it reaches _site_segment /
    # url_path_id, which already reject an empty id with a clear error.
    normalized = _normalize_relative_path(file_path)
    if site_id is not None:
        site_segment = _site_segment(site_id)
        # A "hostname:/server-relative-path" site_id needs a second,
        # closing colon before appending another resource segment, to
        # transition Graph's parser back from path-based site addressing to
        # resource-based addressing (confirmed against Graph's own
        # documented example: ".../sites/contoso.sharepoint.com:/teams/hr:
        # /drive") -- the other two site_id shapes ("root" and the
        # composite "hostname,spSiteId,spWebId" form) never contain ':' and
        # are unaffected.
        site_suffix = ":" if ":" in site_segment else ""
        drive_base = (
            f"/sites/{site_segment}{site_suffix}/drives/{url_path_id(drive_id, 'drive_id')}"
            if drive_id is not None
            else f"/sites/{site_segment}{site_suffix}/drive"
        )
    elif drive_id is not None:
        drive_base = f"/drives/{url_path_id(drive_id, 'drive_id')}"
    else:
        drive_base = "/me/drive"
    return f"{drive_base}/root:/{quote(normalized, safe='/')}:"


def _content_path(file_path: str, site_id: str | None, drive_id: str | None) -> str:
    return f"{_item_path(file_path, site_id, drive_id)}/content"


def _require_etag(value: Any, field_name: str = "expected_etag") -> str:
    if not isinstance(value, str) or not value.strip():
        raise ValueError(f"{field_name} is required")
    normalized = value.strip()
    if "\r" in normalized or "\n" in normalized:
        raise ValueError(f"{field_name} must not contain line breaks")
    return normalized


def _presentation_metadata(
    file_path: str, site_id: str | None, drive_id: str | None
) -> dict[str, Any]:
    item = _graph_request(
        "GET",
        _item_path(file_path, site_id, drive_id),
        params={"$select": "id,size,eTag,@microsoft.graph.downloadUrl"},
    )
    if not isinstance(item, dict) or not item.get("id"):
        raise RuntimeError("Graph did not return presentation metadata")
    size = item.get("size")
    if not isinstance(size, int) or isinstance(size, bool) or size < 0:
        raise RuntimeError("Graph returned an invalid presentation size")
    if size > _MAX_PRESENTATION_BYTES:
        raise ValueError(
            f"The presentation is {size} bytes, over the "
            f"{_MAX_PRESENTATION_BYTES // 1_000_000} MB limit this tool supports"
        )
    _require_etag(item.get("eTag"), "Graph presentation eTag")
    download_url = item.get("@microsoft.graph.downloadUrl")
    if download_url is not None and (
        not isinstance(download_url, str) or not download_url
    ):
        raise RuntimeError("Graph returned an invalid presentation download URL")
    return item


def _download_preauthenticated_content(download_url: str, expected_size: int) -> bytes:
    """Download a signed Graph URL without exposing it or buffering past limits."""
    try:
        response = requests.get(
            download_url,
            stream=True,
            timeout=_BINARY_TIMEOUT_SECONDS,
        )
    except requests.RequestException:
        raise RuntimeError("PowerPoint presentation download failed") from None

    try:
        try:
            response.raise_for_status()
        except requests.HTTPError:
            raise _GraphRequestError(
                "PowerPoint presentation download failed with HTTP "
                f"{response.status_code}",
                status_code=response.status_code,
            ) from None

        content = bytearray()
        try:
            for chunk in response.iter_content(chunk_size=_DOWNLOAD_CHUNK_BYTES):
                if not chunk:
                    continue
                if len(content) + len(chunk) > _MAX_PRESENTATION_BYTES:
                    raise ValueError(
                        "The presentation download exceeded the "
                        f"{_MAX_PRESENTATION_BYTES // 1_000_000} MB limit"
                    )
                content.extend(chunk)
        except requests.RequestException:
            raise RuntimeError("PowerPoint presentation download failed") from None
    finally:
        response.close()

    if len(content) != expected_size:
        raise RuntimeError(
            "PowerPoint presentation size changed while it was being downloaded"
        )
    return bytes(content)


def _download_authenticated_content(
    content_path: str, expected_size: int
) -> bytes:
    """Download Graph ``/content`` when metadata has no signed URL.

    Personal OneDrive and some Graph-compatible drives omit
    ``@microsoft.graph.downloadUrl`` even though the authenticated content
    endpoint is available. Keep the same bounded streaming and safe-error
    behavior as the signed-URL path. ``requests`` follows the normal Graph
    redirect and strips the bearer header when the redirect crosses hosts.
    """
    try:
        response = requests.request(
            method="GET",
            url=f"{GRAPH_BASE_URL}{content_path}",
            headers=_graph_headers(),
            timeout=_BINARY_TIMEOUT_SECONDS,
            stream=True,
        )
    except requests.RequestException:
        raise RuntimeError("PowerPoint presentation download failed") from None

    try:
        try:
            response.raise_for_status()
        except requests.HTTPError:
            raise _GraphRequestError(
                "PowerPoint presentation download failed with HTTP "
                f"{response.status_code}",
                status_code=response.status_code,
            ) from None

        content = bytearray()
        try:
            for chunk in response.iter_content(chunk_size=_DOWNLOAD_CHUNK_BYTES):
                if not chunk:
                    continue
                if len(content) + len(chunk) > _MAX_PRESENTATION_BYTES:
                    raise ValueError(
                        "The presentation download exceeded the "
                        f"{_MAX_PRESENTATION_BYTES // 1_000_000} MB limit"
                    )
                content.extend(chunk)
        except requests.RequestException:
            raise RuntimeError("PowerPoint presentation download failed") from None
    finally:
        response.close()

    if len(content) != expected_size:
        raise RuntimeError(
            "PowerPoint presentation size changed while it was being downloaded"
        )
    return bytes(content)


def _validate_presentation_archive(content: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            infos = archive.infolist()
            if len(infos) > _MAX_PRESENTATION_PARTS:
                raise ValueError(
                    "The presentation contains too many OOXML parts to process safely"
                )
            expanded_size = sum(info.file_size for info in infos)
            if expanded_size > _MAX_PRESENTATION_UNCOMPRESSED_BYTES:
                raise ValueError(
                    "The presentation expands beyond the safe OOXML processing limit"
                )
    except zipfile.BadZipFile as exc:
        raise ValueError(
            "The file is not a valid PowerPoint presentation OOXML archive"
        ) from exc


def _download_presentation(
    file_path: str,
    site_id: str | None,
    drive_id: str | None,
    *,
    expected_etag: str | None = None,
) -> _PresentationSnapshot:
    item = _presentation_metadata(file_path, site_id, drive_id)
    etag = _require_etag(item["eTag"], "Graph presentation eTag")
    if expected_etag is not None and etag != _require_etag(expected_etag):
        raise _ConflictError(
            "The presentation changed after it was read; fetch it again before editing"
        )
    download_url = item.get("@microsoft.graph.downloadUrl")
    if isinstance(download_url, str) and download_url:
        content = _download_preauthenticated_content(download_url, item["size"])
    else:
        content = _download_authenticated_content(
            _content_path(file_path, site_id, drive_id), item["size"]
        )
    _validate_presentation_archive(content)
    try:
        presentation = Presentation(io.BytesIO(content))
    except Exception as exc:
        raise ValueError(
            f"Could not open {file_path!r} as a PowerPoint presentation -- it may "
            "not be a valid .pptx file"
        ) from exc
    return _PresentationSnapshot(presentation=presentation, etag=etag)


def _upload_presentation(
    presentation: PresentationType,
    file_path: str,
    site_id: str | None,
    drive_id: str | None,
    expected_etag: str,
) -> dict[str, Any]:
    buffer = io.BytesIO()
    presentation.save(buffer)
    content = buffer.getvalue()
    if len(content) > _MAX_PRESENTATION_BYTES:
        raise ValueError(
            f"The updated presentation is {len(content)} bytes, over the "
            f"{_MAX_PRESENTATION_BYTES // 1_000_000} MB limit this tool currently "
            "supports"
        )
    # Graph's simple content PUT does not document conditional writes. Use an
    # upload session for every replacement so a stale snapshot is rejected at
    # session creation with If-Match. Graph does not document whether that
    # precondition remains version-fenced through final-fragment commit; see
    # _upload_presentation_session's docstring for the remaining boundary.
    return _upload_presentation_session(
        content, file_path, site_id, drive_id, _require_etag(expected_etag)
    )


def _upload_presentation_session(
    content: bytes,
    file_path: str,
    site_id: str | None,
    drive_id: str | None,
    expected_etag: str,
) -> dict[str, Any]:
    """Replace a presentation through a conditionally created upload session.

    Content is already resident in memory because python-pptx cannot stream a
    save. Progress follows Graph's nextExpectedRanges so a partially accepted
    fragment resumes from the server-reported offset. ``If-Match`` rejects a
    stale version when the session is created. Microsoft Graph does not publish
    a cross-provider commit-time conditional-write contract: deferred source-URL
    commit is Personal-only, while Business/SharePoint use a different final
    POST. Consequently this function must not claim that a change made after
    session creation is proven to be fenced by the final fragment.
    """
    try:
        session = _graph_request(
            "POST",
            f"{_item_path(file_path, site_id, drive_id)}/createUploadSession",
            body={"item": {"@microsoft.graph.conflictBehavior": "replace"}},
            extra_headers={"If-Match": _require_etag(expected_etag)},
        )
    except _GraphRequestError as exc:
        if exc.status_code == 412:
            raise _ConflictError(
                "The presentation changed before the update could be committed"
            ) from None
        raise
    upload_url = session.get("uploadUrl") if isinstance(session, dict) else None
    if not isinstance(upload_url, str) or not upload_url:
        raise RuntimeError("Graph did not return an upload session URL")

    total = len(content)
    result: Any = None
    # The upload-session URL is itself pre-authenticated (a token in its
    # own query string) -- as with _create_only_upload's identical use of
    # this API, this goes through a plain requests.Session() rather than
    # _graph_request (which always attaches an Authorization header, which
    # Graph's docs warn can itself cause a 401 here), and every exception
    # is re-raised "from None" rather than "from exc" so a future
    # traceback/log/APM capture can never surface this URL as __cause__.
    with requests.Session() as http:
        start = 0
        while start < total:
            end = min(start + _UPLOAD_SESSION_CHUNK_BYTES, total)
            is_final_local_fragment = end == total
            try:
                response = http.put(
                    upload_url,
                    data=content[start:end],
                    headers={
                        "Content-Range": f"bytes {start}-{end - 1}/{total}",
                        "Content-Type": _POWERPOINT_MIME_TYPE,
                    },
                    timeout=_BINARY_TIMEOUT_SECONDS,
                )
                response.raise_for_status()
            except requests.HTTPError:
                if response.status_code == 412:
                    _cancel_upload_session(http, upload_url)
                    raise _ConflictError(
                        "The presentation changed before the update could be committed"
                    ) from None
                if is_final_local_fragment and response.status_code >= 500:
                    raise _IndeterminateWriteError(
                        "Graph may have committed the PowerPoint update, but its "
                        "final response was lost; do not retry without reading the "
                        "presentation again"
                    ) from None
                _cancel_upload_session(http, upload_url)
                raise _GraphRequestError(
                    "PowerPoint presentation upload failed with HTTP "
                    f"{response.status_code}",
                    status_code=response.status_code,
                ) from None
            except requests.RequestException:
                if is_final_local_fragment:
                    raise _IndeterminateWriteError(
                        "Graph may have committed the PowerPoint update, but its "
                        "final response was lost; do not retry without reading the "
                        "presentation again"
                    ) from None
                _cancel_upload_session(http, upload_url)
                raise RuntimeError("PowerPoint presentation upload failed") from None

            if response.status_code in (200, 201):
                try:
                    result = response.json()
                except ValueError:
                    raise _IndeterminateWriteError(
                        "Graph accepted the final PowerPoint upload but did not return "
                        "a confirmation item; read the presentation before retrying"
                    ) from None
                break

            try:
                payload = response.json()
                ranges = payload["nextExpectedRanges"]
                offsets = [int(value.split("-", 1)[0]) for value in ranges]
            except (ValueError, TypeError, KeyError, AttributeError) as exc:
                _cancel_upload_session(http, upload_url)
                raise RuntimeError(
                    "Graph returned invalid PowerPoint upload-session progress"
                ) from exc
            if not offsets:
                _cancel_upload_session(http, upload_url)
                raise RuntimeError(
                    "Graph returned empty PowerPoint upload-session progress"
                )
            next_start = min(offsets)
            if not start < next_start <= end:
                _cancel_upload_session(http, upload_url)
                raise RuntimeError(
                    "Graph returned inconsistent PowerPoint upload-session progress"
                )
            start = next_start

    if not isinstance(result, dict) or not result.get("id"):
        raise _IndeterminateWriteError(
            "Graph did not confirm whether the PowerPoint update completed; read the "
            "presentation before retrying"
        )
    safe_item = dict(result)
    safe_item.pop("@microsoft.graph.downloadUrl", None)
    return safe_item


def _cancel_upload_session(http: Any, upload_url: str) -> None:
    """Best-effort cleanup for a session known not to have committed.

    The URL is a bearer secret, so cancellation failures are deliberately
    swallowed without logging or exception chaining.
    """
    try:
        http.delete(upload_url, timeout=_UPLOAD_CANCEL_TIMEOUT_SECONDS)
    except requests.RequestException:
        pass


def _create_only_upload(
    content: bytes, file_path: str, site_id: str | None, drive_id: str | None
) -> dict[str, Any]:
    """Upload content, failing instead of silently overwriting if a file
    already exists at file_path.

    A separate exists-check GET followed by a plain _upload_presentation PUT
    would still race: Graph's own "Upload small files" reference documents
    no If-None-Match or other conditional-write option on the simple
    content PUT, so two concurrent callers can both pass the check and the
    second would silently clobber the first. The upload-session API is
    documented to accept a conflictBehavior of "fail" on session creation,
    which Graph enforces atomically server-side (a 409 nameAlreadyExists if
    the target already exists) -- used here even though the content (a
    blank new presentation) is always small enough for a single-PUT
    session, specifically for that atomicity guarantee.
    """
    item_path = _item_path(file_path, site_id, drive_id)
    try:
        session = _graph_request(
            "POST",
            f"{item_path}/createUploadSession",
            body={"item": {"@microsoft.graph.conflictBehavior": "fail"}},
        )
    except _GraphRequestError as exc:
        if exc.status_code == 409:
            raise ValueError(
                f"{file_path!r} already exists; use the other powerpoint_* "
                "tools to edit it instead of recreating it"
            ) from exc
        raise
    upload_url = session.get("uploadUrl") if isinstance(session, dict) else None
    if not isinstance(upload_url, str) or not upload_url:
        raise RuntimeError("Graph did not return an upload session URL")

    # The upload-session URL is itself pre-authenticated (a token in its own
    # query string) -- Graph's createUploadSession docs warn that including
    # an Authorization header on this PUT can cause a 401 -- so this goes
    # through a plain requests.put, not _graph_request (which always
    # attaches one). Both requests.put's own exception (a connection failure
    # or timeout) and the HTTPError from raise_for_status() embed the full
    # request URL in their default str(), so neither is stringified into a
    # message below, and each is re-raised with "from None" rather than
    # "from exc" -- chaining the original would still attach it as
    # __cause__, which a future traceback/log/APM capture could surface --
    # matching onedrive.py's identical guard on the same hazard.
    try:
        response = requests.put(
            upload_url,
            data=content,
            headers={
                "Content-Range": f"bytes 0-{len(content) - 1}/{len(content)}",
                "Content-Type": _POWERPOINT_MIME_TYPE,
            },
            timeout=_BINARY_TIMEOUT_SECONDS,
        )
        response.raise_for_status()
    except requests.HTTPError:
        status_code = response.status_code
        if status_code == 409:
            _cancel_upload_session(requests, upload_url)
            raise ValueError(
                f"{file_path!r} already exists; use the other powerpoint_* "
                "tools to edit it instead of recreating it"
            ) from None
        if status_code >= 500:
            raise _IndeterminateWriteError(
                "Graph may have created the PowerPoint presentation, but its final "
                "response was lost; check whether the file exists before retrying"
            ) from None
        _cancel_upload_session(requests, upload_url)
        raise _GraphRequestError(
            f"PowerPoint presentation upload failed with HTTP {status_code}",
            status_code=status_code,
        ) from None
    except requests.RequestException:
        raise _IndeterminateWriteError(
            "Graph may have created the PowerPoint presentation, but its final "
            "response was lost; check whether the file exists before retrying"
        ) from None
    try:
        result = response.json()
    except ValueError:
        result = None
    if not isinstance(result, dict) or not result.get("id"):
        raise _IndeterminateWriteError(
            "Graph did not confirm whether the PowerPoint presentation was created; "
            "check whether the file exists before retrying"
        )
    safe_item = dict(result)
    safe_item.pop("@microsoft.graph.downloadUrl", None)
    return safe_item


def _shape_text(shape: Any) -> str | None:
    return shape.text_frame.text if shape.has_text_frame else None


def _slide_texts(slide: Any) -> list[str]:
    """All shape text on slide, recursing into any group and pulling each
    table cell's text.

    slide.shapes alone only yields top-level shapes, and neither a group
    (it has no text frame of its own -- the text lives on the shapes
    nested inside it) nor a table (a GraphicFrame, whose has_text_frame is
    also always False -- the text lives on its individual cells) would
    otherwise contribute anything, so a plain has_text_frame filter over
    slide.shapes silently omits both.
    """
    texts: list[str] = []
    for shape in iter_pptx_shapes(slide.shapes):
        text = _shape_text(shape)
        if text:
            texts.append(text)
        elif getattr(shape, "has_table", False):
            for row in shape.table.rows:
                for cell in row.cells:
                    text = cell.text_frame.text
                    if text:
                        texts.append(text)
    return texts


def _slide_notes(slide: Any) -> str | None:
    """slide's speaker notes text, or None if it has no notes slide or no
    notes placeholder on it.

    Checked via has_notes_slide rather than a bare getattr/try -- python-
    pptx's own notes_slide property *creates* a notes slide on first access
    if one doesn't already exist, which would silently mutate a
    presentation this module never intends to write back for a read-only
    tool. notes_text_frame can still be None even when has_notes_slide is
    True -- per its own docstring, that happens if the notes placeholder
    was deleted from the notes slide (the notes-slide XML part survives,
    just without a body placeholder).
    """
    if not slide.has_notes_slide:
        return None
    notes_text_frame = slide.notes_slide.notes_text_frame
    if notes_text_frame is None:
        return None
    return notes_text_frame.text or None


def _select_body_placeholder(slide: Any) -> Any | None:
    """Find the first non-title placeholder on slide that can hold text.

    A placeholder's idx alone doesn't guarantee it has a text frame -- a
    picture, chart, or other content placeholder can sit at a lower idx
    than a real text placeholder (verified: python-pptx's own
    SlidePlaceholders iterates in idx order, with no guarantee idx order
    matches "text placeholders first"), so this skips non-text
    placeholders entirely rather than stopping at the first non-title one
    regardless of whether it can actually hold text.
    """
    return next(
        (
            ph
            for ph in slide.placeholders
            if ph.placeholder_format.idx != 0 and ph.has_text_frame
        ),
        None,
    )


def _require_int(value: Any, field_name: str) -> int:
    """Validate an integer-typed tool argument.

    FastMCP/Pydantic validates and coerces arguments against a tool's type
    hints for a real call over the MCP protocol, but that validation is
    bypassed by a caller that invokes this Python function directly (e.g. a
    test, or any other in-process caller) -- without this, a non-int index
    reaches python-pptx's own indexing/comparison and fails with a raw,
    less clear TypeError instead. bool is excluded even though it's a
    subclass of int in Python, since True/False are never valid indices.
    """
    if not isinstance(value, int) or isinstance(value, bool):
        raise TypeError(f"{field_name} must be an integer")
    return value


def _text_frame_has_dynamic_content(text_frame: Any) -> bool:
    """Whether text_frame holds a hyperlink or a dynamic field (e.g. an
    auto-updating slide number or date placeholder) that a full-frame text
    replacement would silently discard.

    python-pptx's TextFrame.text setter clears every existing paragraph
    and rebuilds the frame from a single new run per line (verified
    directly: assigning .text removes all <a:p> children) -- a run's
    <a:hlinkClick> and a paragraph-level <a:fld> element are both lost
    with no error or warning.
    """
    tx_body = text_frame._txBody
    return (
        tx_body.find(".//" + qn("a:hlinkClick")) is not None
        or tx_body.find(".//" + qn("a:hlinkMouseOver")) is not None
        or tx_body.find(".//" + qn("a:fld")) is not None
    )


def _text_frame_has_distinct_run_formatting(text_frame: Any) -> bool:
    """Whether replacement would collapse genuinely distinct run formatting.

    Multiple adjacent runs are common even when their direct formatting is
    identical. Collapsing those runs is safe because replacement preserves the
    first run's ``a:rPr``. Remain conservative when the direct XML differs: it
    can carry properties python-pptx does not expose through ``Font``.
    """
    for paragraph in text_frame.paragraphs:
        formatting: set[str | None] = set()
        for run in paragraph.runs:
            run_properties = run._r.find(qn("a:rPr"))
            formatting.add(
                str(run_properties.xml) if run_properties is not None else None
            )
        # A DrawingML line break may carry direct character properties even
        # though python-pptx does not expose it through paragraph.runs. An
        # absent br/rPr inherits surrounding formatting and is therefore not a
        # distinct style; an explicit one must participate in the guard.
        for line_break in paragraph._p.findall(qn("a:br")):
            run_properties = line_break.find(qn("a:rPr"))
            if run_properties is not None:
                formatting.add(str(run_properties.xml))
        if len(formatting) > 1:
            return True
    return False


def _replace_text_frame_text(text_frame: Any, text: str) -> None:
    """Replace text_frame's text, preserving each existing paragraph's
    paragraph-level formatting (alignment, bullet/numbering, indent level)
    and its first run's character formatting (bold, italic, font, size,
    color) by position.

    TextFrame.text's own setter (python-pptx) clears every <a:p> and
    rebuilds one new paragraph per "\n"-separated line, with additional
    runs around a "\v" soft break. Those generated elements have no
    <a:pPr>/<a:rPr> at all. Existing paragraphs keep their formatting by
    position; additional paragraphs inherit the final old paragraph's
    formatting, and every generated run in a paragraph receives that old
    paragraph's first-run formatting, including the generated ``a:br`` node.
    Paragraph-end character properties are retained by position as well.
    Removed paragraphs have no output counterpart and are discarded.
    """
    saved: list[tuple[Any, Any, Any]] = []
    for paragraph in text_frame.paragraphs:
        p_pr = paragraph._p.find(qn("a:pPr"))
        first_run = paragraph.runs[0] if paragraph.runs else None
        r_pr = first_run._r.find(qn("a:rPr")) if first_run is not None else None
        if r_pr is None:
            first_break = paragraph._p.find(qn("a:br"))
            r_pr = first_break.find(qn("a:rPr")) if first_break is not None else None
        end_para_r_pr = paragraph._p.find(qn("a:endParaRPr"))
        if r_pr is None:
            r_pr = end_para_r_pr
        saved.append(
            (
                deepcopy(p_pr) if p_pr is not None else None,
                deepcopy(r_pr) if r_pr is not None else None,
                deepcopy(end_para_r_pr) if end_para_r_pr is not None else None,
            )
        )

    text_frame.text = text

    for index, paragraph in enumerate(text_frame.paragraphs):
        p_pr, r_pr, end_para_r_pr = saved[min(index, len(saved) - 1)]
        if p_pr is not None:
            existing_p_pr = paragraph._p.find(qn("a:pPr"))
            if existing_p_pr is not None:
                paragraph._p.remove(existing_p_pr)
            paragraph._p.insert(0, deepcopy(p_pr))
        if r_pr is not None:
            for run in paragraph.runs:
                existing_r_pr = run._r.find(qn("a:rPr"))
                if existing_r_pr is not None:
                    run._r.remove(existing_r_pr)
                run._r.insert(0, deepcopy(r_pr))
            for line_break in paragraph._p.findall(qn("a:br")):
                existing_r_pr = line_break.find(qn("a:rPr"))
                if existing_r_pr is not None:
                    line_break.remove(existing_r_pr)
                line_break.insert(0, deepcopy(r_pr))
        if end_para_r_pr is not None:
            existing_end_para_r_pr = paragraph._p.find(qn("a:endParaRPr"))
            if existing_end_para_r_pr is not None:
                paragraph._p.remove(existing_end_para_r_pr)
            paragraph._p.append(deepcopy(end_para_r_pr))


def _require_slide(presentation: PresentationType, slide_index: int) -> Any:
    slides = presentation.slides
    if not 0 <= slide_index < len(slides):
        raise ValueError(
            f"slide_index {slide_index} is out of range for a presentation with "
            f"{len(slides)} slides"
        )
    return slides[slide_index]


def _delete_slide(presentation: PresentationType, slide_index: int) -> None:
    """Remove a slide by index.

    python-pptx has no public API for removing a slide (confirmed: no
    Slides.remove()/delete() method exists in the library). The documented
    community workaround -- dropping the slide's <p:sldId> entry from the
    presentation's slide-id list, which is what actually determines slide
    order/membership in the underlying XML -- is used here instead.

    Removing only the <p:sldId> entry leaves the slide's own XML part and
    the presentation-to-slide relationship in the saved package -- verified
    directly: python-pptx's serializer does no orphan-pruning on save, so a
    caller reading the package's parts/relationships directly (rather than
    walking presentation.slides) would still see the "deleted" slide's
    content, and repeated add/delete cycles would accumulate dead parts.
    part.drop_rel() below removes that relationship (and, when no supported
    inbound reference remains, the part itself) so the deleted slide's content
    doesn't survive in the saved file at all. Custom shows, section metadata,
    and hyperlinks from another slide can reference the target independently;
    those cases fail closed because silently leaving a dangling reference or
    rewriting the user's navigation structure would both be unsafe.
    """
    slide_id_list = presentation.slides._sldIdLst
    slide_ids = list(slide_id_list)
    if not 0 <= slide_index < len(slide_ids):
        raise ValueError(
            f"slide_index {slide_index} is out of range for a presentation with "
            f"{len(slide_ids)} slides"
        )
    target = slide_ids[slide_index]
    relationship_id = target.get(qn("r:id"))
    target_slide_id = target.get("id")
    if relationship_id:
        for element in presentation._element.iter():
            if element is target:
                continue
            if element.get(qn("r:id")) == relationship_id:
                raise ValueError(
                    "The slide is referenced by a custom slide show or other "
                    "presentation feature and cannot be deleted safely"
                )
            local_name = (
                element.tag.rsplit("}", 1)[-1] if isinstance(element.tag, str) else ""
            )
            if (
                local_name == "sldId"
                and target_slide_id is not None
                and element.get("id") == target_slide_id
            ):
                raise ValueError(
                    "The slide is referenced by a presentation section and cannot "
                    "be deleted safely"
                )

        target_part = presentation.part.related_part(relationship_id)
        for other_slide in presentation.slides:
            if other_slide.part is target_part:
                continue
            for relationship in other_slide.part.rels.values():
                if (
                    not relationship.is_external
                    and relationship.target_part is target_part
                ):
                    raise ValueError(
                        "The slide is linked from another slide and cannot be "
                        "deleted safely"
                    )

    slide_id_list.remove(target)
    if relationship_id:
        presentation.part.drop_rel(relationship_id)


@mcp.tool()
def powerpoint_create_presentation(
    file_path: str, site_id: str | None = None, drive_id: str | None = None
) -> str:
    """Create a new, blank PowerPoint presentation at file_path. Fails if a
    file already exists there -- edit it with the other powerpoint_* tools
    instead of recreating it."""
    try:
        buffer = io.BytesIO()
        Presentation().save(buffer)
        item = _create_only_upload(buffer.getvalue(), file_path, site_id, drive_id)
        return _success(item=item)
    except _IndeterminateWriteError as e:
        return _indeterminate(str(e))
    except Exception as e:
        logger.error("Error creating PowerPoint presentation %s: %s", file_path, e)
        return _error(str(e))


@mcp.tool()
def powerpoint_get_presentation_text(
    file_path: str,
    site_id: str | None = None,
    drive_id: str | None = None,
    cursor: str | None = None,
) -> str:
    """Get all text in a PowerPoint presentation, as a list of slides each
    with the text of every text-bearing shape on it (including a shape
    nested inside a group, and each cell of a table) plus that slide's
    speaker notes, if any. Large decks are returned as bounded pages; pass
    next_cursor back as cursor to continue. The cursor is tied to the current
    eTag, so pagination fails with conflict if the presentation changes. A
    single slide that cannot fit is reported as omitted_item and the returned
    cursor advances to the following slide."""
    try:
        snapshot = _download_presentation(file_path, site_id, drive_id)
        presentation = snapshot.presentation
        total_count = len(presentation.slides)
        scope = "presentation_text"
        start_index = _decode_read_cursor(
            cursor, etag=snapshot.etag, scope=scope, total_count=total_count
        )

        def slide_at(slide_index: int) -> dict[str, Any]:
            slide = presentation.slides[slide_index]
            return {
                "slide_index": slide_index,
                "shapes": _slide_texts(slide),
                "notes": _slide_notes(slide),
            }

        return _bounded_page_response(
            field_name="slides",
            total_count=total_count,
            start_index=start_index,
            item_at=slide_at,
            item_label="slide",
            etag=snapshot.etag,
            scope=scope,
        )
    except _ConflictError as e:
        return _conflict(str(e))
    except Exception as e:
        logger.error(
            "Error getting text for PowerPoint presentation %s: %s", file_path, e
        )
        return _error(str(e))


@mcp.tool()
def powerpoint_list_slides(
    file_path: str,
    site_id: str | None = None,
    drive_id: str | None = None,
    cursor: str | None = None,
) -> str:
    """List a PowerPoint presentation's slides, with each slide's index,
    layout name, and shape count. Returns an etag that must be supplied to
    mutating tools so a stale positional request is rejected before its upload
    session is created. Pass next_cursor back as cursor when truncated is
    true."""
    try:
        snapshot = _download_presentation(file_path, site_id, drive_id)
        presentation = snapshot.presentation
        total_count = len(presentation.slides)
        scope = "slide_list"
        start_index = _decode_read_cursor(
            cursor, etag=snapshot.etag, scope=scope, total_count=total_count
        )

        def slide_at(slide_index: int) -> dict[str, Any]:
            slide = presentation.slides[slide_index]
            return {
                "slide_index": slide_index,
                "layout_name": slide.slide_layout.name,
                "shape_count": len(slide.shapes),
            }

        return _bounded_page_response(
            field_name="slides",
            total_count=total_count,
            start_index=start_index,
            item_at=slide_at,
            item_label="slide",
            etag=snapshot.etag,
            scope=scope,
        )
    except _ConflictError as e:
        return _conflict(str(e))
    except Exception as e:
        logger.error(
            "Error listing slides for PowerPoint presentation %s: %s", file_path, e
        )
        return _error(str(e))


@mcp.tool()
def powerpoint_get_slide_text(
    file_path: str,
    slide_index: int,
    site_id: str | None = None,
    drive_id: str | None = None,
    cursor: str | None = None,
) -> str:
    """Get one slide's shapes with their index, type, and text (needed by
    powerpoint_set_shape_text). Returns an etag that must be passed to that
    tool together with the positional indices. Large shape lists are returned
    as bounded pages; pass next_cursor back as cursor to continue. A single
    shape that cannot fit is reported as omitted_item and the returned cursor
    advances to the following shape."""
    try:
        slide_index = _require_int(slide_index, "slide_index")
        snapshot = _download_presentation(file_path, site_id, drive_id)
        presentation = snapshot.presentation
        slide = _require_slide(presentation, slide_index)
        total_count = len(slide.shapes)
        scope = f"slide_text:{slide_index}"
        start_index = _decode_read_cursor(
            cursor, etag=snapshot.etag, scope=scope, total_count=total_count
        )

        def shape_at(shape_index: int) -> dict[str, Any]:
            shape = slide.shapes[shape_index]
            return {
                "shape_index": shape_index,
                "shape_type": str(shape.shape_type),
                "is_placeholder": shape.is_placeholder,
                "text": _shape_text(shape),
            }

        return _bounded_page_response(
            field_name="shapes",
            total_count=total_count,
            start_index=start_index,
            item_at=shape_at,
            item_label="shape",
            etag=snapshot.etag,
            scope=scope,
        )
    except _ConflictError as e:
        return _conflict(str(e))
    except Exception as e:
        logger.error(
            "Error getting slide %s text for PowerPoint presentation %s: %s",
            slide_index,
            file_path,
            e,
        )
        return _error(str(e))


@mcp.tool()
def powerpoint_set_shape_text(
    file_path: str,
    slide_index: int,
    shape_index: int,
    text: str,
    expected_etag: str,
    site_id: str | None = None,
    drive_id: str | None = None,
) -> str:
    """Replace a shape's text on a slide, by slide_index and shape_index
    (from powerpoint_get_slide_text), together with that tool's etag as
    expected_etag. Only works on a shape that already
    has a text frame (a title, body, or plain text box); errors otherwise.
    Preserves each existing paragraph's alignment/bullet/indent formatting
    and its single run's character formatting (bold, italic, font, size,
    color) by position. Additional paragraphs (a "\\n" in text starts a new
    paragraph) inherit the final existing paragraph's formatting, and runs
    and the break node generated by a "\\v" soft break inherit their
    paragraph's character formatting; paragraph-end character properties are
    retained too. Errors rather than silently discarding multiple differently
    formatted runs or breaks, hyperlinks, or dynamic fields such as an
    auto-updating slide number or date; edit those shapes directly in
    PowerPoint instead."""
    try:
        slide_index = _require_int(slide_index, "slide_index")
        shape_index = _require_int(shape_index, "shape_index")
        expected_etag = _require_etag(expected_etag)
        snapshot = _download_presentation(
            file_path, site_id, drive_id, expected_etag=expected_etag
        )
        presentation = snapshot.presentation
        slide = _require_slide(presentation, slide_index)
        shape_count = len(slide.shapes)
        if not 0 <= shape_index < shape_count:
            raise ValueError(
                f"shape_index {shape_index} is out of range for slide {slide_index} "
                f"with {shape_count} shapes"
            )
        shape = slide.shapes[shape_index]
        if not shape.has_text_frame:
            raise ValueError(
                f"shape {shape_index} on slide {slide_index} has no text frame "
                "(e.g. it's an image or a plain line/connector) and cannot hold text"
            )
        if _text_frame_has_dynamic_content(shape.text_frame):
            raise ValueError(
                f"shape {shape_index} on slide {slide_index} contains a hyperlink "
                "or a dynamic field (e.g. slide number or date) that replacing its "
                "text would silently delete -- edit this shape directly in "
                "PowerPoint instead"
            )
        if _text_frame_has_distinct_run_formatting(shape.text_frame):
            raise ValueError(
                f"shape {shape_index} on slide {slide_index} contains multiple text "
                "runs whose formatting cannot be preserved by whole-frame "
                "replacement -- edit this shape directly in PowerPoint instead"
            )
        _replace_text_frame_text(shape.text_frame, text)
        item = _upload_presentation(
            presentation, file_path, site_id, drive_id, expected_etag
        )
        return _success(item=item)
    except _ConflictError as e:
        return _conflict(str(e))
    except _IndeterminateWriteError as e:
        return _indeterminate(str(e))
    except Exception as e:
        logger.error(
            "Error setting shape %s text on slide %s in PowerPoint presentation %s: %s",
            shape_index,
            slide_index,
            file_path,
            e,
        )
        return _error(str(e))


@mcp.tool()
def powerpoint_add_slide(
    file_path: str,
    expected_etag: str,
    title: str | None = None,
    body_text: str | None = None,
    layout_index: int = _DEFAULT_SLIDE_LAYOUT_INDEX,
    site_id: str | None = None,
    drive_id: str | None = None,
) -> str:
    """Add a new slide at the end of a PowerPoint presentation.

    expected_etag must come from powerpoint_list_slides or
    powerpoint_list_slide_layouts. layout_index selects a slide layout from
    the presentation's slide master (0 is usually "Title Slide", 1 "Title
    and Content", 6 "Blank" --
    use powerpoint_list_slide_layouts to see what this presentation
    actually has, since layouts vary by template). title, if given, is set
    on the layout's title placeholder when present. body_text, if given, is
    set on the first non-title placeholder (by idx order) that has a text
    frame -- a non-text placeholder earlier in idx order (e.g. a picture or
    chart placeholder) is skipped rather than causing body_text to be left
    unset. This is still a heuristic, not a guaranteed "body" role: on a
    "Title Slide" layout, for example, that placeholder is actually the
    subtitle. Use powerpoint_get_slide_text after adding the slide to
    confirm where each piece of text actually landed."""
    try:
        layout_index = _require_int(layout_index, "layout_index")
        expected_etag = _require_etag(expected_etag)
        snapshot = _download_presentation(
            file_path, site_id, drive_id, expected_etag=expected_etag
        )
        presentation = snapshot.presentation
        layouts = presentation.slide_layouts
        if not 0 <= layout_index < len(layouts):
            raise ValueError(
                f"layout_index {layout_index} is out of range for this "
                f"presentation's {len(layouts)} slide layouts"
            )
        slide = presentation.slides.add_slide(layouts[layout_index])
        title_applied = False
        if title is not None:
            if slide.shapes.title is None:
                raise ValueError(
                    f"slide layout {layout_index} has no title placeholder; "
                    "title was not applied"
                )
            slide.shapes.title.text = title
            title_applied = True
        body_applied = False
        if body_text is not None:
            body_placeholder = _select_body_placeholder(slide)
            if body_placeholder is None:
                raise ValueError(
                    f"slide layout {layout_index} has no body text placeholder; "
                    "body_text was not applied"
                )
            body_placeholder.text_frame.text = body_text
            body_applied = True
        item = _upload_presentation(
            presentation, file_path, site_id, drive_id, expected_etag
        )
        return _success(
            item=item,
            slide_index=len(presentation.slides) - 1,
            title_applied=title_applied,
            body_applied=body_applied,
        )
    except _ConflictError as e:
        return _conflict(str(e))
    except _IndeterminateWriteError as e:
        return _indeterminate(str(e))
    except Exception as e:
        logger.error(
            "Error adding slide to PowerPoint presentation %s: %s", file_path, e
        )
        return _error(str(e))


@mcp.tool()
def powerpoint_list_slide_layouts(
    file_path: str,
    site_id: str | None = None,
    drive_id: str | None = None,
    cursor: str | None = None,
) -> str:
    """List the slide layouts available in a presentation's template, with
    the index powerpoint_add_slide's layout_index expects, plus the etag that
    tool requires to reject a stale snapshot before creating its upload
    session. Pass next_cursor back as cursor when truncated is true."""
    try:
        snapshot = _download_presentation(file_path, site_id, drive_id)
        presentation = snapshot.presentation
        total_count = len(presentation.slide_layouts)
        scope = "slide_layouts"
        start_index = _decode_read_cursor(
            cursor, etag=snapshot.etag, scope=scope, total_count=total_count
        )

        def layout_at(layout_index: int) -> dict[str, Any]:
            return {
                "layout_index": layout_index,
                "name": presentation.slide_layouts[layout_index].name,
            }

        return _bounded_page_response(
            field_name="layouts",
            total_count=total_count,
            start_index=start_index,
            item_at=layout_at,
            item_label="layout",
            etag=snapshot.etag,
            scope=scope,
        )
    except _ConflictError as e:
        return _conflict(str(e))
    except Exception as e:
        logger.error(
            "Error listing slide layouts for PowerPoint presentation %s: %s",
            file_path,
            e,
        )
        return _error(str(e))


@mcp.tool()
def powerpoint_delete_slide(
    file_path: str,
    slide_index: int,
    expected_etag: str,
    site_id: str | None = None,
    drive_id: str | None = None,
) -> str:
    """Delete a slide by index from the snapshot identified by expected_etag."""
    try:
        slide_index = _require_int(slide_index, "slide_index")
        expected_etag = _require_etag(expected_etag)
        snapshot = _download_presentation(
            file_path, site_id, drive_id, expected_etag=expected_etag
        )
        presentation = snapshot.presentation
        _delete_slide(presentation, slide_index)
        item = _upload_presentation(
            presentation, file_path, site_id, drive_id, expected_etag
        )
        return _success(item=item)
    except _ConflictError as e:
        return _conflict(str(e))
    except _IndeterminateWriteError as e:
        return _indeterminate(str(e))
    except Exception as e:
        logger.error(
            "Error deleting slide %s from PowerPoint presentation %s: %s",
            slide_index,
            file_path,
            e,
        )
        return _error(str(e))


if __name__ == "__main__":
    mcp.run()
